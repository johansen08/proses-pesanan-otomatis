"""Uji proses_label.py dengan server Jubelio tiruan (tanpa akses internet).

Jalankan:  .venv\\Scripts\\python tests\\test_proses_label.py
Bagian yang memakai HTML label asli hanya jalan jika rekaman sniff 092943 masih ada.
"""
import json
import re
import sys
import tempfile
from pathlib import Path
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import proses_label as pl  # noqa: E402

REKAMAN = ROOT / "sniff" / "sniff_output" / "sniff_jubel_20260926_092943.har"
SKU = "T01-BSBI-5"


class Resp:
    def __init__(self, status=200, data=None, content=None, headers=None):
        self.status_code = status
        self._data = data
        self.content = content if content is not None else json.dumps(data).encode()
        self.text = self.content.decode("utf-8", "ignore")
        self.headers = headers or {"content-type": "application/json"}

    def json(self):
        return self._data


class JubelioPalsu:
    """Meniru perilaku yang terlihat di rekaman 26-09-2026."""

    def __init__(self, html_label: str, tolak_sekali=True):
        self.html_label = html_label
        self.tolak_sekali = tolak_sekali
        self.log = []                                    # (method, path, body)
        self.orders = [
            {"salesorder_id": 9068214, "salesorder_no": "TT-A", "shipper": "J&T Express Hemat", "grand_total": "35900.0000", "total_qty": "1.0000"},
            {"salesorder_id": 9068180, "salesorder_no": "TT-B", "shipper": "J&T Express NEXT-DAY DELIVERY", "grand_total": "35900.0000", "total_qty": "1.0000"},
            {"salesorder_id": 9068161, "salesorder_no": "TT-C", "shipper": "J&T Express Hemat", "grand_total": "35900.0000", "total_qty": "1.0000"},
            {"salesorder_id": 9067835, "salesorder_no": "SP-D", "shipper": "SPX Hemat", "grand_total": "24042.0000", "total_qty": "1.0000"},
            {"salesorder_id": 9068999, "salesorder_no": "SP-KREATOR", "shipper": "SPX Hemat", "grand_total": "0.0000", "total_qty": "1.0000"},
            {"salesorder_id": 9067659, "salesorder_no": "TT-SUDAH-DIPAKAI", "shipper": "J&T Express Hemat", "grand_total": "35900.0000", "total_qty": "1.0000"},
            {"salesorder_id": 9060000, "salesorder_no": "TT-BUKAN-SPESIAL", "shipper": "J&T Express Hemat", "grand_total": "35900.0000", "total_qty": "1.0000"},
        ]
        self.picklist = None
        self.cek_picklist_setelah_selesai = 0
        self.cek_finish_pick = 0
        self.panggil_resi = 0
        self.info_dok = {}
        self.batal_ids = set()
        self.batal_hanya_di_detail = False      # respons resi belum ikut berubah status

    # ---------------------------------------------------------------- util
    def _catat(self, method, url, body):
        self.log.append((method, urlsplit(url).path, body))

    def get(self, url, params=None, headers=None, timeout=None, cookies=None):
        self._catat("GET", url, params)
        path = urlsplit(url).path
        if "report-prod" in url and path == "/":
            assert cookies == {"JB_OMNI_ACCESS_TOKEN": "TKN"}
            return Resp(content=self.html_label.encode(), headers={"content-type": "text/html"})
        if path.endswith("/info"):
            doc = path.split("/")[-2]
            self.info_dok[doc] = self.info_dok.get(doc, 0) + 1
            return Resp(202 if self.info_dok[doc] < 3 else 200, {})
        if "/documents/" in path:
            return Resp(content=b"%PDF-1.4 label palsu", headers={"content-type": "application/pdf"})
        assert headers and headers.get("authorization") == "TKN", path
        if path.endswith("ready-to-process/"):
            return Resp(data={"data": self.orders, "totalCount": len(self.orders)})
        m = re.search(r"sales/picklists/(\d+)$", path)
        if m:
            p = self.picklist
            if p["is_completed"]:
                self.cek_picklist_setelah_selesai += 1
                if self.cek_picklist_setelah_selesai >= 2:      # butuh waktu (merah -> hitam)
                    for i in p["items"]:
                        i["wms_status"], i["qty_picked"], i["bin_id"] = "FINISH_PICK", "1.0000", 14
            data = json.loads(json.dumps(p))
            data["is_completed"] = p["is_completed"] and self.cek_picklist_setelah_selesai >= 2
            return Resp(data=data)
        if "default-bin" in path:
            return Resp(data={"bin_id": 14, "location_id": -1, "bin_final_code": "LX-BX-KX-RX-1"})
        if path.endswith("finish-pick/"):
            self.cek_finish_pick += 1
            if not (self.picklist and self.picklist["is_completed"]) or self.cek_finish_pick < 2:
                return Resp(data={"data": [], "totalCount": 0})
            ids = {i["salesorder_id"] for i in self.picklist["items"]}
            data = [dict(o, picklist_no=self.picklist["picklist_no"]) for o in self.orders
                    if o["salesorder_id"] in ids]
            return Resp(data={"data": data, "totalCount": len(data)})
        if path.endswith("shipping-label/"):
            return Resp(data={"status": "ok", "url": "https://report-prod.jubelio.com/?&token=RPT",
                              "title": "Label Pengiriman"})
        m = re.search(r"sales/orders/(\d+)$", path)
        if m:
            batal = int(m.group(1)) in self.batal_ids
            return Resp(data={"salesorder_id": int(m.group(1)), "channel_status": "CANCELLED" if batal else
                              "AWAITING_SHIPMENT", "internal_status": "CANCELED" if batal else "PROCESSING"})
        raise AssertionError(f"GET tak dikenal {url}")

    def post(self, url, json=None, headers=None, timeout=None, cookies=None):
        self._catat("POST", url, json)
        path = urlsplit(url).path
        if "report-prod" in url:
            assert cookies == {"JB_OMNI_ACCESS_TOKEN": "TKN"}
            if path.endswith("/clients"):
                return Resp(data={"clientId": "c1"})
            if path.endswith("/parameters"):
                return Resp(data=[{"id": k, "value": v} for k, v in json["parameterValues"].items()])
            if path.endswith("/instances"):
                return Resp(201, {"instanceId": "i1"})
            if path.endswith("/documents"):
                return Resp(202, {"documentId": f"d{len(self.info_dok) + 1}"})
        assert headers.get("authorization") == "TKN", path
        if path.endswith("items-to-pick/"):
            return Resp(data=[{"salesorder_detail_id": 14500000 + i, "item_id": 9436, "location_id": -1,
                               "qty_ordered": "1.0000", "salesorder_id": i, "bundle_item_id": 0,
                               "package_detail_id": None, "package_id": None, "end_qty": "90.0000",
                               "item_full_name": f"{SKU} - Topi Baseball", "salesorder_no": "x"}
                              for i in json["ids"]])
        if path.endswith("wms/sales/picklists/"):
            if not json["is_completed"]:
                if 9067659 in json["salesorderIds"]:
                    if self.tolak_sekali:
                        # setelah ditolak, pesanan itu hilang dari Siap Proses (seperti di rekaman)
                        self.orders = [o for o in self.orders if o["salesorder_id"] != 9067659]
                    return Resp(500, {"statusCode": 500, "message": "An internal server error occurred",
                                      "code": "error: Pesanan sudah dipakai di transaksi lain. Pesanan: TT-SUDAH-DIPAKAI"})
                self.picklist = {"picklist_id": 154839, "picklist_no": "PICK-000154839", "note": None,
                                 "is_completed": False,
                                 "items": [{"picklist_detail_id": 13590000 + n, "item_id": x["item_id"],
                                            "location_id": -1, "qty_ordered": "1.0000", "qty_picked": "0.0000",
                                            "salesorder_detail_id": x["salesorder_detail_id"],
                                            "salesorder_id": x["salesorder_id"], "bundle_item_id": 0,
                                            "package_id": None, "package_detail_id": 0, "invoice_no": None,
                                            "wms_status": "PICK", "bin_id": None, "item_code": SKU}
                                           for n, x in enumerate(json["items"])]}
                self.orders = [o for o in self.orders if o["salesorder_id"] not in json["salesorderIds"]] + \
                    [o for o in self.orders if o["salesorder_id"] in json["salesorderIds"]]
            else:
                self.picklist["is_completed"] = True
            return Resp(data={"status": "ok", "data": {"picks": [
                {"picklist_id": 154839, "picklist_no": "PICK-000154839", "status": "ok"}], "invalidSO": []}})
        if path.endswith("shipper-pickup-time/"):
            return Resp(data=[{"shipper": "SPX Hemat", "timeSlots": []}])
        if path.endswith("shipments/orders/"):
            self.panggil_resi += 1
            return Resp(data=[{"salesorder_id": i, "salesorder_no": f"SO{i}",
                               "internal_status": "CANCELED" if i in self.batal_ids
                               and not self.batal_hanya_di_detail else "PROCESSING",
                               "tracking_no": f"JY{i}" if self.panggil_resi >= 3 and i not in self.batal_ids
                               else ""} for i in json["ids"]])
        raise AssertionError(f"POST tak dikenal {url}")


def _html_label() -> str:
    if REKAMAN.exists():
        har = json.loads(REKAMAN.read_text(encoding="utf-8"))
        for e in har["log"]["entries"]:
            t = e["response"]["content"].get("text") or ""
            if e["request"]["url"].startswith("https://report-prod.jubelio.com/?") and "Label Pengiriman" in t[:5000]:
                return t
    return ('<script>jQuery("#reportViewer").telerik_ReportViewer({"serviceUrl":"/api/reports/",'
            '"reportSource":{"report":"Label Pengiriman-x-3227","parameters":{"ids":["1"],'
            '"list":"[{\\"a\\":\\"});\\"}]"}},"viewMode":"PRINT_PREVIEW"});</script>')


def uji_report_source_sama_dengan_rekaman():
    if not REKAMAN.exists():
        print("  (lewati: rekaman tidak ada)")
        return
    har = json.loads(REKAMAN.read_text(encoding="utf-8"))
    E = har["log"]["entries"]
    param_req = [json.loads(e["request"]["postData"]["text"]) for e in E
                 if e["request"]["url"].endswith("/parameters") and e["request"].get("postData")]
    n = 0
    for e in E:
        if e["request"]["url"].startswith("https://report-prod.jubelio.com/?"):
            rs = pl._report_source(e["response"]["content"]["text"])
            assert {"report": rs["report"], "parameterValues": rs["parameters"]} in param_req
            n += 1
    assert n >= 2
    print(f"  {n} halaman report asli terbaca, sama persis dengan body /parameters yang dikirim web")


def uji_mode_uji_tidak_mengubah_apapun():
    j = JubelioPalsu(_html_label())
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    resi = {"TT-A", "TT-B", "TT-C", "SP-D", "SP-KREATOR", "TT-SUDAH-DIPAKAI"}
    hasil = pl.rencana(k, {SKU: sorted(resi)})
    assert all(m == "GET" for m, _, _ in j.log), j.log
    assert hasil[0]["pesanan"] == 5                     # kreator & bukan-spesial dibuang
    print("  mode uji: hanya GET, 5 pesanan akan diproses (kreator & bukan spesial dibuang)")


def uji_proses_lengkap():
    j = JubelioPalsu(_html_label())
    tidur = []
    k = pl.Klien("TKN", sesi=j, tidur=tidur.append)
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        resi = sorted({"TT-A", "TT-B", "TT-C", "SP-D", "SP-KREATOR", "TT-SUDAH-DIPAKAI"})
        hasil = pl.proses(k, {SKU: resi}, d / "label", d / "riwayat.xlsx", {SKU: "1B-B2-2"})

        assert len(hasil) == 1 and hasil[0]["No Picklist"] == "PICK-000154839", hasil
        assert hasil[0]["Rak"] == "1B-B2-2" and hasil[0]["Durasi"].endswith("detik"), hasil
        assert hasil[0]["Total Pesanan"] == 4 and hasil[0]["Resi Keluar"] == 4, hasil
        pdf = Path(hasil[0]["File Label"])
        assert pdf.exists() and pdf.read_bytes().startswith(b"%PDF")
        assert pdf.name.startswith(f"PICK-000154839_SPESIAL_{SKU}_"), \
            f"nama file picklist SKU spesial harus memuat penanda SPESIAL: {pdf.name}"
        assert pdf.parent.name == pl.TAG_SPESIAL and pdf.parent.parent == d / "label", \
            f"label SKU spesial harus disimpan di subfolder {pl.TAG_SPESIAL}: {pdf}"

        posts = [(p, b) for m, p, b in j.log if m == "POST" and "core-api" in p]
        buat = [b for p, b in posts if p.endswith("wms/sales/picklists/") and not b["is_completed"]]
        assert len(buat) == 2 and 9067659 in buat[0]["salesorderIds"], "harus coba ulang setelah 500"
        assert sorted(buat[1]["salesorderIds"]) == [9067835, 9068161, 9068180, 9068214]
        assert set(buat[1]) == {"is_completed", "is_warehouse", "items", "merge_location", "picker_id",
                                "picklist_id", "picklist_no", "salesorderIds"}
        assert set(buat[1]["items"][0]) == {"salesorder_detail_id", "item_id", "location_id", "qty_ordered",
                                            "salesorder_id", "bundle_item_id", "package_detail_id", "package_id"}
        assert buat[1]["items"][0]["qty_ordered"] == 1 and buat[1]["items"][0]["package_id"] == 0

        selesai = [b for p, b in posts if p.endswith("wms/sales/picklists/") and b["is_completed"]]
        assert len(selesai) == 1
        assert set(selesai[0]) == {"is_completed", "is_warehouse", "items", "picklist_id", "picklist_no", "note"}
        assert set(selesai[0]["items"][0]) == {
            "bin_id", "bundle_item_id", "item_id", "location_id", "picklist_detail_id", "qty_ordered",
            "qty_picked", "salesorder_detail_id", "salesorder_id", "invoice_no", "update", "package_id",
            "package_detail_id"}
        assert selesai[0]["items"][0]["bin_id"] == 14 and selesai[0]["items"][0]["qty_picked"] == 1

        # urutan: selesaikan -> tunggu FINISH_PICK -> Picking>Selesai -> resi -> label
        urut = [p.rsplit("core-api/", 1)[-1] for m, p, _ in j.log if "core-api" in p]
        i_selesai = max(i for i, p in enumerate(urut) if p == "wms/sales/picklists/")
        assert "wms/sales/v2/orders/finish-pick/" in urut[i_selesai:]
        assert j.cek_picklist_setelah_selesai >= 2, "harus menunggu status FINISH_PICK"
        assert j.panggil_resi == 3, "harus mengulang minta resi sampai terisi"
        assert JEDA_RESI in tidur
        i_label = urut.index("reports/shipping-label/")
        assert i_label > max(i for i, p in enumerate(urut) if p == "wms/sales/shipments/orders/")

        from openpyxl import load_workbook
        ws = load_workbook(d / "riwayat.xlsx").active
        baris = list(ws.values)
        assert baris[0][:4] == ("Waktu", "SKU", "No Picklist", "Total Pesanan")
        assert baris[0][-1] == "Durasi" and baris[1][-1] == hasil[0]["Durasi"], baris
        assert baris[1][1:5] == (SKU, "PICK-000154839", 4, 4), baris[1]
    print("  proses lengkap: 500 -> coba ulang, tunggu FINISH_PICK, 3x minta resi, PDF & riwayat tersimpan")


def uji_lanjutkan_picklist_terhenti():
    j = JubelioPalsu(_html_label())
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    j.picklist = {"picklist_id": 154839, "picklist_no": "PICK-000154839", "note": None, "is_completed": True,
                  "items": [{"picklist_detail_id": 1, "item_id": 9436, "location_id": -1, "qty_ordered": "1.0000",
                             "salesorder_detail_id": 5, "salesorder_id": 9068214, "bundle_item_id": 0,
                             "wms_status": "FINISH_PICK", "item_code": SKU}]}
    j.cek_picklist_setelah_selesai = 5
    with tempfile.TemporaryDirectory() as d:
        baris = pl.lanjutkan(k, "PICK-000154839", Path(d) / "label", Path(d) / "riwayat.xlsx")
    assert not any(m == "POST" and p.endswith("wms/sales/picklists/") for m, p, _ in j.log), \
        "picking yang sudah selesai tidak boleh dikirim ulang"
    assert baris["Resi Keluar"] == 1 and baris["File Label"]
    print("  lanjutkan: picklist sudah FINISH_PICK tidak diselesaikan ulang, langsung resi & label")


def _picklist_selesai(j, ids):
    j.picklist = {"picklist_id": 154839, "picklist_no": "PICK-000154839", "note": None, "is_completed": True,
                  "items": [{"picklist_detail_id": n, "item_id": 9436, "location_id": -1, "qty_ordered": "1.0000",
                             "salesorder_detail_id": n, "salesorder_id": i, "bundle_item_id": 0,
                             "wms_status": "FINISH_PICK", "item_code": SKU} for n, i in enumerate(ids)]}
    j.cek_picklist_setelah_selesai = 5


def uji_pesanan_batal_tidak_ditunggu():
    j = JubelioPalsu(_html_label())
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    _picklist_selesai(j, [9068214, 9068180])
    j.batal_ids = {9068180}
    with tempfile.TemporaryDirectory() as d:
        baris = pl.lanjutkan(k, "PICK-000154839", Path(d) / "label", Path(d) / "riwayat.xlsx")
    assert j.panggil_resi == 3, "pesanan batal tidak boleh ditunggu sampai batas waktu"
    assert baris["Resi Keluar"] == 1 and baris["File Label"], baris
    assert baris["Catatan"] == "batal, tidak dicetak: SO9068180", baris
    label = [b for m, p, b in j.log if p.endswith("shipping-label/")]
    assert label == [{"ids[0]": 9068214, "tz": "Asia/Jakarta"}], label
    print("  pesanan batal: tidak ditunggu, sisanya tetap dicetak, dihitung selesai")


def uji_pesanan_batal_terlihat_di_detail():
    j = JubelioPalsu(_html_label())
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    _picklist_selesai(j, [9068214, 9068180])
    j.batal_ids, j.batal_hanya_di_detail = {9068180}, True
    j.panggil_resi = 2                                    # resi TT-A keluar di panggilan berikutnya
    lama, pl.TUNGGU_RESI_S = pl.TUNGGU_RESI_S, 0
    try:
        with tempfile.TemporaryDirectory() as d:
            baris = pl.lanjutkan(k, "PICK-000154839", Path(d) / "label", Path(d) / "riwayat.xlsx")
    finally:
        pl.TUNGGU_RESI_S = lama
    assert any(p.endswith("sales/orders/9068180") for m, p, _ in j.log), "harus cek detail pesanan"
    assert baris["Resi Keluar"] == 1 and baris["Catatan"] == "batal, tidak dicetak: SO9068180", baris
    print("  status batal hanya di detail pesanan: tetap dikenali setelah batas waktu")


def uji_pesanan_tanpa_resi_dicatat_peringatan():
    """Pesanan yang bukan batal tapi tidak kunjung dapat resi (mis. ada request cancel yang
    masih diproses) harus dicatat lewat peringatan_resi, bukan cuma masuk log biasa."""
    j = JubelioPalsu(_html_label())
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    _picklist_selesai(j, [9068214, 9068180])
    lama, pl.TUNGGU_RESI_S = pl.TUNGGU_RESI_S, 0
    pl.peringatan_resi._sesi.clear()
    try:
        with tempfile.TemporaryDirectory() as d:
            baris = pl.lanjutkan(k, "PICK-000154839", Path(d) / "label", Path(d) / "riwayat.xlsx")
    finally:
        pl.TUNGGU_RESI_S = lama
    assert baris["Catatan"] == "belum dapat resi: SO9068214, SO9068180", baris
    assert len(pl.peringatan_resi._sesi) == 1
    pesan = pl.peringatan_resi._sesi[0]
    assert "SO9068214" in pesan and "SO9068180" in pesan and "PICK-000154839" in pesan, pesan
    print("  pesanan tanpa resi (bukan batal): dicatat peringatan_resi untuk diinfokan ke CS")


def uji_lewati_jika_kurang_dari_3():
    j = JubelioPalsu(_html_label())
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    with tempfile.TemporaryDirectory() as d:
        hasil = pl.proses(k, {SKU: ["TT-A", "TT-B"]}, Path(d), Path(d) / "r.xlsx")
        assert not (Path(d) / "r.xlsx").exists(), "SKU dilewati tidak dicatat di riwayat"
    assert len(hasil) == 1 and hasil[0]["Catatan"].startswith("Dilewati"), hasil
    assert "No Picklist" not in hasil[0] and hasil[0]["Durasi"]
    assert not any(m == "POST" for m, _, _ in j.log)
    print("  < 3 pesanan: SKU dilewati tanpa membuat picklist")


class JubelioPalsuUrgent:
    """Server tiruan khusus skenario picklist urgent (lintas SKU)."""

    def __init__(self):
        self.log = []
        # 3 Lazada asli (source 4) + kebocoran channel lain yang sengaja diselipkan untuk
        # memastikan ambil_pesanan_channel menyaring ulang, bukan cuma percaya query API
        self.pesanan_lazada = [
            {"salesorder_id": 100 + i, "source": 4, "shipper": "J&T", "grand_total": "35900.0000"}
            for i in range(3)
        ] + [{"salesorder_id": 999, "source": 64, "shipper": "SPX",
              "grand_total": "35900.0000"}]     # Shopee nyelip
        # 200 Tokopedia asli (source 128) + 50 TikTok "Shop | Tokopedia" (source 131076) -
        # keduanya harus IKUT karena skenario ini tidak difilter channel, cuma kurir
        self.pesanan_gtl_sicepat = [
            {"salesorder_id": 200 + i, "source": 128, "shipper": "GTL", "grand_total": "35900.0000"}
            for i in range(200)
        ] + [{"salesorder_id": 400 + i, "source": 131076, "shipper": "SiCepat",
              "grand_total": "35900.0000"} for i in range(50)]
        self.picklist_no = 0
        self.stok_kosong_ids: set[int] = set()

    def get(self, url, params=None, headers=None, timeout=None, cookies=None):
        self.log.append(("GET", url, params))
        assert headers.get("authorization") == "TKN"
        if url.endswith("ready-to-process/"):
            assert params.get("sort_by") == "transaction_date" and params.get("sort_direction") == "ASC", \
                "resi terlama harus diambil duluan, supaya masuk picklist pertama kalau dipecah"
            assert "order_type[0]" not in params, \
                "Lazada/GTL-SiCepat urgent tidak pakai SPX atau Shopee -> jangan pakai filter tipe"
            if params.get("channel_ids[0]") == 4:
                data = self.pesanan_lazada
            else:
                assert "channel_ids[0]" not in params, \
                    "GTL/SiCepat urgent lintas channel -> jangan difilter channel_ids"
                assert params.get("couriers[0]") == "gtl" and params.get("couriers[1]") == "sicepat"
                data = self.pesanan_gtl_sicepat
            return Resp(data={"data": data, "totalCount": len(data)})
        raise AssertionError(f"GET tak dikenal {url}")

    def post(self, url, json=None, headers=None, timeout=None, cookies=None):
        self.log.append(("POST", url, json))
        assert headers.get("authorization") == "TKN"
        if url.endswith("items-to-pick/"):
            return Resp(data=[{"salesorder_detail_id": 9000 + i, "item_id": 1, "location_id": -1,
                               "qty_ordered": "1.0000", "salesorder_id": i, "bundle_item_id": 0,
                               "package_detail_id": None, "package_id": None, "end_qty": "999.0000",
                               "item_full_name": "X", "salesorder_no": f"SO{i}"} for i in json["ids"]])
        if url.endswith("wms/sales/empty-stock"):
            self.stok_kosong_ids |= set(json["salesorder_ids"])
            return Resp(data={"status": "ok"})
        if url.endswith("wms/sales/picklists/"):
            self.picklist_no += 1
            no = f"PICK-{900000 + self.picklist_no}"
            return Resp(data={"status": "ok", "data": {
                "picks": [{"picklist_id": self.picklist_no, "picklist_no": no, "status": "ok"}],
                "invalidSO": []}})
        raise AssertionError(f"POST tak dikenal {url}")


def uji_ambil_pesanan_channel_resi_terlama_masuk_batch_pertama():
    # Terbukti dari sniff 29-09-2026 15:47 (sort_direction=ASC di web Jubelio -> transaction_date
    # menaik/terlama dulu): kalau totalnya > 200 dan dipecah bagi_batch(), resi terlama harus
    # selalu ada di batch pertama (picklist pertama), bukan tertahan di batch belakangan.
    class JubelioPalsuUrut:
        def get(self, url, params=None, headers=None, timeout=None, cookies=None):
            assert params.get("sort_direction") == "ASC"
            # API sudah mengurutkan ASC di server; klien tinggal percaya urutan responsnya
            data = [{"salesorder_id": i, "source": pl.CHANNEL_ID_LAZADA, "grand_total": "35900.0000",
                     "transaction_date": f"2026-09-29T{6 + i // 40:02d}:00:00Z"}
                    for i in range(250)]     # id 0 = terlama (jam 06), id 249 = terbaru
            return Resp(data={"data": data, "totalCount": len(data)})

    k = pl.Klien("TKN", sesi=JubelioPalsuUrut(), tidur=lambda s: None)
    pesanan = pl.ambil_pesanan_channel(k, [pl.CHANNEL_ID_LAZADA])
    batch = pl.bagi_batch([o["salesorder_id"] for o in pesanan])
    assert batch[0] == list(range(200)), "200 resi terlama harus di picklist pertama"
    assert batch[1] == list(range(200, 250)), "50 resi terbaru baru di picklist kedua"
    print("  250 resi (id 0 = terlama) -> batch pertama isinya id 0-199 (terlama), "
          "bukan tercampur/id terbaru duluan")


def uji_ambil_pesanan_channel_nilai_0_hanya_dibuang_untuk_tiktok_shop():
    # TT-586350230929114342-67824 (01-10-2026): pesanan TikTok Shop nilai 0 lolos sampai
    # finish-pick karena ambil_pesanan_channel() belum memfilternya. Aturan dikonfirmasi user:
    # nilai 0/kosong cuma berarti sampel/kreator utk channel TikTok Shop ("Shop | Tokopedia",
    # source 131076) - Shopee/Lazada/channel lain nilai kecil (bahkan 0) tetap pesanan
    # sungguhan, jangan ikut dibuang.
    class JubelioPalsuCampur:
        def get(self, url, params=None, headers=None, timeout=None, cookies=None):
            data = [
                {"salesorder_id": 1, "salesorder_no": "TT-SAMPEL", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": "0.0000"},
                {"salesorder_id": 2, "salesorder_no": "TT-SAMPEL-KOSONG", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": None},
                {"salesorder_id": 3, "salesorder_no": "TT-ASLI", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": "35900.0000"},
                {"salesorder_id": 4, "salesorder_no": "SP-NILAI-0", "source": pl.CHANNEL_ID_SHOPEE, "grand_total": "0.0000"},
                {"salesorder_id": 5, "salesorder_no": "LZ-NILAI-0", "source": pl.CHANNEL_ID_LAZADA, "grand_total": "0.0000"},
            ]
            return Resp(data={"data": data, "totalCount": len(data)})

    k = pl.Klien("TKN", sesi=JubelioPalsuCampur(), tidur=lambda s: None)
    pesanan = pl.ambil_pesanan_channel(k)
    assert {o["salesorder_no"] for o in pesanan} == {"TT-ASLI", "SP-NILAI-0", "LZ-NILAI-0"}, pesanan
    print("  TT-SAMPEL & TT-SAMPEL-KOSONG (TikTok Shop, nilai 0/kosong) dibuang; SP-NILAI-0 & "
          "LZ-NILAI-0 (channel lain, nilai 0) TETAP diproses - bukan sampel")


def uji_ambil_pesanan_sampel_hanya_ambil_tiktok_shop_nilai_0():
    # Kebalikan dari ambil_pesanan_channel(): ambil_pesanan_sampel() HARUS cuma mengambil
    # pesanan yang dibuang di sana (TikTok Shop nilai 0/kosong), supaya pesanan itu sekarang
    # masuk picklist sampel-nya sendiri, bukan hilang sama sekali.
    class JubelioPalsuCampur:
        def get(self, url, params=None, headers=None, timeout=None, cookies=None):
            assert params.get("channel_ids[0]") == pl.CHANNEL_ID_TIKTOK_SHOP
            data = [
                {"salesorder_id": 1, "salesorder_no": "TT-SAMPEL", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": "0.0000"},
                {"salesorder_id": 2, "salesorder_no": "TT-SAMPEL-KOSONG", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": None},
                {"salesorder_id": 3, "salesorder_no": "TT-ASLI", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": "35900.0000"},
            ]
            return Resp(data={"data": data, "totalCount": len(data)})

    k = pl.Klien("TKN", sesi=JubelioPalsuCampur(), tidur=lambda s: None)
    pesanan = pl.ambil_pesanan_sampel(k)
    assert {o["salesorder_no"] for o in pesanan} == {"TT-SAMPEL", "TT-SAMPEL-KOSONG"}, pesanan
    print("  ambil_pesanan_sampel(): hanya TT-SAMPEL & TT-SAMPEL-KOSONG, TT-ASLI (nilai asli) tidak ikut")


def uji_proses_sampel_buat_1_picklist_dan_dilewati_kalau_kosong():
    class JubelioPalsuSampel:
        def get(self, url, params=None, headers=None, timeout=None, cookies=None):
            data = [
                {"salesorder_id": 1, "salesorder_no": "TT-SAMPEL-1", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": "0.0000"},
                {"salesorder_id": 2, "salesorder_no": "TT-SAMPEL-2", "source": pl.CHANNEL_ID_TIKTOK_SHOP, "grand_total": "0.0000"},
            ]
            return Resp(data={"data": data, "totalCount": len(data)})

        def post(self, url, json=None, headers=None, timeout=None, cookies=None):
            if url.endswith("items-to-pick/"):
                return Resp(data=[{"salesorder_detail_id": 9000 + i, "item_id": 1, "location_id": -1,
                                   "qty_ordered": "1.0000", "salesorder_id": i, "bundle_item_id": 0,
                                   "package_detail_id": None, "package_id": None, "end_qty": "999.0000",
                                   "item_full_name": "X", "salesorder_no": f"SO{i}"} for i in json["ids"]])
            if url.endswith("wms/sales/picklists/"):
                return Resp(data={"status": "ok", "data": {
                    "picks": [{"picklist_id": 1, "picklist_no": "PICK-000900001", "status": "ok"}],
                    "invalidSO": []}})
            raise AssertionError(f"POST tak dikenal {url}")

    panggilan = []
    asli = pl.lanjutkan_picklist

    def stub(k, pid, pno, jumlah, sku, folder_label, nama_file=None):
        panggilan.append((pid, pno, jumlah, sku))
        return {"Waktu": "-", "SKU": sku, "No Picklist": pno, "Total Pesanan": jumlah,
                "Resi Keluar": jumlah, "File Label": f"{pno}_{sku}_x.pdf", "Catatan": ""}
    pl.lanjutkan_picklist = stub
    try:
        k = pl.Klien("TKN", sesi=JubelioPalsuSampel(), tidur=lambda s: None)
        with tempfile.TemporaryDirectory() as d:
            riwayat = Path(d) / "riwayat.xlsx"
            hasil = pl.proses_sampel(k, riwayat, Path(d) / "label")
    finally:
        pl.lanjutkan_picklist = asli

    assert len(hasil) == 1 and hasil[0]["SKU"] == pl.LABEL_SAMPEL, hasil
    assert panggilan and panggilan[0][3] == pl.LABEL_SAMPEL, panggilan
    print("  proses_sampel(): 2 pesanan TikTok Shop nilai 0 -> 1 picklist SAMPEL-TIKTOK")

    class JubelioPalsuKosong:
        def get(self, url, params=None, headers=None, timeout=None, cookies=None):
            return Resp(data={"data": [], "totalCount": 0})

    k = pl.Klien("TKN", sesi=JubelioPalsuKosong(), tidur=lambda s: None)
    with tempfile.TemporaryDirectory() as d:
        hasil = pl.proses_sampel(k, Path(d) / "riwayat.xlsx", Path(d) / "label")
    assert hasil == [], hasil
    print("  proses_sampel(): tidak ada pesanan sampel -> dilewati, tidak ada picklist dibuat")


def uji_urgent_menyaring_channel_bocor_dan_membagi_batch():
    j = JubelioPalsuUrgent()
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)

    lazada = pl.ambil_pesanan_channel(k, [pl.CHANNEL_ID_LAZADA])
    assert [o["salesorder_id"] for o in lazada] == [100, 101, 102], lazada
    print("  Lazada: pesanan channel lain (Shopee) ikut kefilter API tapi dibuang lagi di kita")

    gtl_sicepat = pl.ambil_pesanan_channel(k, couriers=pl.KURIR_FILTER_URGENT_GTL_SICEPAT)
    assert len(gtl_sicepat) == 250
    assert {128, 131076} == {o["source"] for o in gtl_sicepat}
    print('  GTL/SiCepat: tidak difilter channel -> Tokopedia asli DAN "Shop | Tokopedia" '
          "(TikTok) sama-sama ikut, urgent-nya ditentukan kurir")

    # lanjutkan_picklist() (selesaikan picking -> resi -> label) sudah diuji lengkap lewat
    # uji_proses_lengkap(); di sini cukup pastikan proses_urgent() memanggilnya dengan label
    # yang benar (dipakai lanjutkan_picklist utk nama file & kolom SKU riwayat).
    panggilan = []
    asli = pl.lanjutkan_picklist

    def stub(k, pid, pno, jumlah, sku, folder_label, nama_file=None):
        panggilan.append((pid, pno, jumlah, sku, folder_label))
        return {"Waktu": "-", "SKU": sku, "No Picklist": pno, "Total Pesanan": jumlah,
                "Resi Keluar": jumlah, "File Label": f"{pno}_{sku}_x.pdf", "Catatan": ""}
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            riwayat = Path(d) / "riwayat.xlsx"
            hasil = pl.proses_urgent(k, riwayat, Path(d) / "label")
            from openpyxl import load_workbook
            baris = list(load_workbook(riwayat).active.values)
    finally:
        pl.lanjutkan_picklist = asli

    assert all(sku in ("LAZADA", "GTL-SICEPAT") for _, _, _, sku, _ in panggilan), panggilan
    per_channel = {}
    for h in hasil:
        per_channel.setdefault(h["SKU"], []).append(h)
    assert len(per_channel["LAZADA"]) == 1
    assert per_channel["LAZADA"][0]["Total Pesanan"] == 3
    assert per_channel["LAZADA"][0]["File Label"] == f"{per_channel['LAZADA'][0]['No Picklist']}_LAZADA_x.pdf"
    assert len(per_channel["GTL-SICEPAT"]) == 2, "250 pesanan -> 2 picklist (200 + 50)"
    assert per_channel["GTL-SICEPAT"][0]["Total Pesanan"] == 200
    assert per_channel["GTL-SICEPAT"][1]["Total Pesanan"] == 50
    assert len(baris) == 1 + 3, "3 picklist urgent tercatat di riwayat"
    print("  proses_urgent: 3 pesanan Lazada -> 1 picklist, 250 GTL/SiCepat -> 2 picklist "
          "(200+50), nama file pakai label skenario (mis. PICK-..._LAZADA_...)")

    # --channel: cuma jalankan 1 skenario
    skenario_lazada = [s for s in pl.SKENARIO_URGENT if s[0] == "Lazada"]
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            hasil = pl.proses_urgent(k, Path(d) / "riwayat.xlsx", Path(d) / "label", skenario_lazada)
    finally:
        pl.lanjutkan_picklist = asli
    assert {h["SKU"] for h in hasil} == {"LAZADA"}, hasil
    print("  skenario bisa dibatasi 1 channel saja (dipakai --channel)")


def uji_urgent_jam_tunda_ditahan_lalu_lanjut_setelah_jam_16():
    import datetime as dt

    # Unit _saring_jam_urgent(): WIB 13:30 (<=14:00) lolos, WIB 14:30 (>14:00) ditahan,
    # tanpa transaction_date tidak pernah ditahan (lebih aman diproses daripada hilang).
    pesanan = [
        {"salesorder_id": 1, "transaction_date": "2026-10-02T06:30:00Z"},   # 13:30 WIB
        {"salesorder_id": 2, "transaction_date": "2026-10-02T07:30:00Z"},   # 14:30 WIB
        {"salesorder_id": 3},
    ]
    sebelum_jam_16 = dt.datetime(2026, 10, 2, 13, 0, 0, tzinfo=pl.WIB)
    pakai, ditahan = pl._saring_jam_urgent(pesanan, pl.JAM_CUTOFF_URGENT_LAZADA, sebelum_jam_16)
    assert ditahan == 1 and {o["salesorder_id"] for o in pakai} == {1, 3}, (pakai, ditahan)
    print("  Lazada jam 13:00: pesanan WIB 14:30 (di atas cutoff 14:00) ditahan, WIB 13:30 & "
          "tanpa transaction_date tetap diproses")

    # Setelah JAM_LANJUT_URGENT (16:00), cutoff diabaikan - semua pesanan (termasuk yang
    # tadinya ditahan) langsung diproses.
    setelah_jam_16 = dt.datetime(2026, 10, 2, 16, 0, 0, tzinfo=pl.WIB)
    pakai2, ditahan2 = pl._saring_jam_urgent(pesanan, pl.JAM_CUTOFF_URGENT_LAZADA, setelah_jam_16)
    assert ditahan2 == 0 and len(pakai2) == 3, (pakai2, ditahan2)
    print("  setelah jam 16:00: batas jam diabaikan, semua pesanan (termasuk yang tadinya "
          "ditahan) diproses")

    # Integrasi: proses_urgent() dengan JubelioPalsuUrgent, salah satu pesanan Lazada &
    # GTL/SiCepat diberi jam pesan di atas cutoff masing-masing.
    j = JubelioPalsuUrgent()
    j.pesanan_lazada[0]["transaction_date"] = "2026-10-02T07:30:00Z"   # 14:30 WIB -> ditahan
    j.pesanan_gtl_sicepat[0]["transaction_date"] = "2026-10-02T08:30:00Z"   # 15:30 WIB -> ditahan
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)

    panggilan = []
    asli = pl.lanjutkan_picklist

    def stub(k, pid, pno, jumlah, sku, folder_label, nama_file=None):
        panggilan.append((pid, pno, jumlah, sku, folder_label))
        return {"Waktu": "-", "SKU": sku, "No Picklist": pno, "Total Pesanan": jumlah,
                "Resi Keluar": jumlah, "File Label": f"{pno}_{sku}_x.pdf", "Catatan": ""}
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            hasil = pl.proses_urgent(k, Path(d) / "riwayat.xlsx", Path(d) / "label",
                                     sekarang=sebelum_jam_16)
    finally:
        pl.lanjutkan_picklist = asli
    per_sku = {h["SKU"]: h for h in hasil if h.get("No Picklist")}
    assert per_sku["LAZADA"]["Total Pesanan"] == 2, per_sku      # 3 - 1 ditahan
    assert sum(h["Total Pesanan"] for h in hasil if h["SKU"] == "GTL-SICEPAT") == 249, hasil
    print("  proses_urgent() jam 13:00: 1 Lazada & 1 GTL/SiCepat ditahan, sisanya tetap jadi picklist")

    # Dipanggil lagi setelah jam 16:00 (mis. siklus TIPE 1 berikutnya) -> yang tadinya
    # ditahan ikut diproses, tanpa batas jam lagi.
    j2 = JubelioPalsuUrgent()
    j2.pesanan_lazada[0]["transaction_date"] = "2026-10-02T07:30:00Z"
    j2.pesanan_gtl_sicepat[0]["transaction_date"] = "2026-10-02T08:30:00Z"
    k2 = pl.Klien("TKN", sesi=j2, tidur=lambda s: None)
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            hasil2 = pl.proses_urgent(k2, Path(d) / "riwayat.xlsx", Path(d) / "label",
                                      sekarang=setelah_jam_16)
    finally:
        pl.lanjutkan_picklist = asli
    assert sum(h["Total Pesanan"] for h in hasil2 if h["SKU"] == "LAZADA") == 3, hasil2
    assert sum(h["Total Pesanan"] for h in hasil2 if h["SKU"] == "GTL-SICEPAT") == 250, hasil2
    print("  proses_urgent() jam 16:00: semua pesanan diproses, termasuk yang tadinya ditahan")


class JubelioPalsuReguler:
    """Server tiruan khusus skenario picklist sisa reguler (TikTok Shop & Shopee)."""

    def __init__(self):
        self.log = []
        # TikTok Shop (source 131076): SO-1..SO-5 total_qty 1 (calon 1 qty reguler),
        # SO-6 total_qty 2 (calon kombinasi). SO-1 & SO-2 "sudah" SKU spesial -> dikeluarkan.
        # + Shopee (source 64): SO-7 total_qty 3 (kombinasi), SO-8 total_qty 1 (1 qty reguler).
        # + kebocoran channel lain (source 4, Lazada) utk pastikan disaring ulang.
        # + SO-10 nilai 0 (sampel/kreator, mis. TT-586350230929114342-67824) -> harus
        # dikeluarkan ambil_pesanan_channel() sebelum sempat masuk pisah_reguler().
        self.pesanan = [
            {"salesorder_id": 1, "salesorder_no": "SO-1", "source": 131076, "total_qty": "1.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 2, "salesorder_no": "SO-2", "source": 131076, "total_qty": "1.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 3, "salesorder_no": "SO-3", "source": 131076, "total_qty": "1.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 4, "salesorder_no": "SO-4", "source": 131076, "total_qty": "1.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 5, "salesorder_no": "SO-5", "source": 131076, "total_qty": "1.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 6, "salesorder_no": "SO-6", "source": 131076, "total_qty": "2.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 7, "salesorder_no": "SO-7", "source": 64, "total_qty": "3.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 8, "salesorder_no": "SO-8", "source": 64, "total_qty": "1.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 9, "salesorder_no": "SO-9", "source": 4, "total_qty": "1.0000", "grand_total": "35900.0000"},
            {"salesorder_id": 10, "salesorder_no": "SO-10", "source": 131076, "total_qty": "1.0000", "grand_total": "0.0000"},
        ]
        self.picklist_no = 0

    def get(self, url, params=None, headers=None, timeout=None, cookies=None):
        self.log.append(("GET", url, params))
        assert headers.get("authorization") == "TKN"
        if url.endswith("ready-to-process/"):
            assert params.get("sort_by") == "transaction_date" and params.get("sort_direction") == "ASC", \
                "resi terlama harus diambil duluan, supaya masuk picklist pertama kalau dipecah"
            assert params.get("channel_ids[0]") == pl.CHANNEL_ID_TIKTOK_SHOP
            assert params.get("channel_ids[1]") == pl.CHANNEL_ID_SHOPEE
            assert params.get("couriers[0]") == "j&t" and params.get("couriers[1]") == "spx"
            assert [params[f"order_type[{i}]"] for i in range(len(pl.TIPE_PESANAN_FILTER))] \
                == pl.TIPE_PESANAN_FILTER, \
                "channel Shopee & kurir SPX -> pesanan kilat harus difilter"
            return Resp(data={"data": self.pesanan, "totalCount": len(self.pesanan)})
        raise AssertionError(f"GET tak dikenal {url}")

    def post(self, url, json=None, headers=None, timeout=None, cookies=None):
        self.log.append(("POST", url, json))
        assert headers.get("authorization") == "TKN"
        if url.endswith("items-to-pick/"):
            return Resp(data=[{"salesorder_detail_id": 9000 + i, "item_id": 1, "location_id": -1,
                               "qty_ordered": "1.0000", "salesorder_id": i, "bundle_item_id": 0,
                               "package_detail_id": None, "package_id": None, "end_qty": "999.0000",
                               "item_full_name": "X", "salesorder_no": f"SO{i}"} for i in json["ids"]])
        if url.endswith("wms/sales/picklists/"):
            self.picklist_no += 1
            no = f"PICK-{910000 + self.picklist_no}"
            return Resp(data={"status": "ok", "data": {
                "picks": [{"picklist_id": self.picklist_no, "picklist_no": no, "status": "ok"}],
                "invalidSO": []}})
        raise AssertionError(f"POST tak dikenal {url}")


class JubelioPalsuKombinasiRak:
    """Server tiruan endpoint zones-racks-combination, 2 halaman (page_size kecil) supaya
    paging ikut teruji."""

    def __init__(self, kombinasi: list[str], page_size: int = 2):
        self.kombinasi = kombinasi
        self.page_size = page_size
        self.log = []

    def get(self, url, params=None, headers=None, timeout=None, cookies=None):
        self.log.append(("GET", url, params))
        assert headers.get("authorization") == "TKN"
        assert url.endswith("zones-racks-combination")
        assert params.get("status") == "PAID"
        assert params.get("combination_type") == "racks"
        page = params["page"]
        awal = (page - 1) * self.page_size
        potongan = self.kombinasi[awal:awal + self.page_size]
        data = [{"combination": c} for c in potongan]
        return Resp(data={"data": data, "totalCount": len(self.kombinasi)})


def uji_ambil_kombinasi_rak_buang_gabungan_dan_kosong():
    kombinasi = ["", "2A-B1-1", "3A-C2-2", "2A-B1-1 - 3B-H1-2", "1B-A2-3"]
    j = JubelioPalsuKombinasiRak(kombinasi, page_size=2)
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    hasil = pl.ambil_kombinasi_rak(k)
    assert hasil == ["2A-B1-1", "3A-C2-2", "1B-A2-3"], hasil
    assert len(j.log) == 3, "5 baris / page_size 2 -> 3 halaman"
    print("  ambil_kombinasi_rak: string kosong & gabungan (\" - \") dibuang, paging jalan")


def uji_kelompokkan_kombinasi_per_grup():
    kombinasi = ["2A-B1-1", "2A-B2-2", "3A-C2-2", "1B-A2-3", "2B-D1-1", "3B-E1-1",
                 "4C-F1-1", "1A-X1-1"]
    hasil = pl.kelompokkan_kombinasi_per_grup(kombinasi)
    assert list(hasil.keys()) == pl.GRUP_RAK, hasil
    assert hasil["2A"] == ["2A-B1-1", "2A-B2-2"]
    assert hasil["3A"] == ["3A-C2-2"]
    assert hasil["1B"] == ["1B-A2-3"]
    assert hasil["2B"] == ["2B-D1-1"]
    assert hasil["3B"] == ["3B-E1-1"]
    print("  kelompokkan_kombinasi_per_grup: prefix di luar GRUP_RAK (4C, 1A) diabaikan, "
          "urutan key = GRUP_RAK")


def uji_reguler_maksimal_200_per_picklist():
    # pisah_reguler() murni fungsi data (tanpa API); bagi_batch() dipakai proses_reguler()
    # lewat _proses_channel_batch() yang sama persis dengan proses_urgent() -> cukup buktikan
    # hasil pisah_reguler() dibagi bagi_batch() dengan aturan 200 yang sama.
    pesanan = [{"salesorder_id": i, "salesorder_no": f"SO-{i}", "total_qty": "1.0000"}
              for i in range(250)] + [{"salesorder_id": 900 + i, "salesorder_no": f"KB-{i}",
                                       "total_qty": "2.0000"} for i in range(210)]
    satu_qty, kombinasi = pl.pisah_reguler(pesanan, resi_spesial_semua=set())
    assert len(satu_qty) == 250 and len(kombinasi) == 210

    batch_1qty = pl.bagi_batch([o["salesorder_id"] for o in satu_qty])
    batch_kombinasi = pl.bagi_batch([o["salesorder_id"] for o in kombinasi])
    assert [len(b) for b in batch_1qty] == [200, 50], batch_1qty
    assert [len(b) for b in batch_kombinasi] == [200, 10], batch_kombinasi
    print("  1QTY-REGULER 250 pesanan -> 200+50, KOMBINASI-REGULER 210 pesanan -> 200+10 "
          "(maks 200/picklist, sisanya dipecah otomatis)")


def uji_reguler_keluarkan_spesial_dan_pisah_1qty_kombinasi():
    j = JubelioPalsuReguler()
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    resi_spesial_semua = {"SO-1", "SO-2"}

    pesanan = pl.ambil_pesanan_reguler(k)
    assert {o["salesorder_no"] for o in pesanan} == {f"SO-{i}" for i in range(1, 9)}, pesanan
    print("  ambil_pesanan_reguler: SO-9 (channel Lazada, bocor dari filter API) dan SO-10 "
          "(nilai 0, sampel/kreator) sama-sama dibuang lagi di kita")
    satu_qty, kombinasi = pl.pisah_reguler(pesanan, resi_spesial_semua)
    assert {o["salesorder_no"] for o in satu_qty} == {"SO-3", "SO-4", "SO-5", "SO-8"}, satu_qty
    assert {o["salesorder_no"] for o in kombinasi} == {"SO-6", "SO-7"}, kombinasi
    print("  pisah_reguler: SO-1/SO-2 (spesial) dikeluarkan, sisanya kepisah 1 qty vs "
          "kombinasi (qty>1) dengan benar")

    panggilan = []
    asli = pl.lanjutkan_picklist

    def stub(k, pid, pno, jumlah, sku, folder_label, nama_file=None):
        panggilan.append((pid, pno, jumlah, sku, folder_label))
        return {"Waktu": "-", "SKU": sku, "No Picklist": pno, "Total Pesanan": jumlah,
                "Resi Keluar": jumlah, "File Label": f"{pno}_{sku}_x.pdf", "Catatan": ""}
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            riwayat = Path(d) / "riwayat.xlsx"
            hasil = pl.proses_reguler(k, resi_spesial_semua, riwayat, Path(d) / "label")
            from openpyxl import load_workbook
            baris = list(load_workbook(riwayat).active.values)
    finally:
        pl.lanjutkan_picklist = asli

    per_label = {h["SKU"]: h for h in hasil}
    assert per_label["1QTY-REGULER"]["Total Pesanan"] == 4
    assert per_label["KOMBINASI-REGULER"]["Total Pesanan"] == 2
    assert len(baris) == 1 + 2, "2 picklist reguler tercatat di riwayat"
    print("  proses_reguler: 1 picklist 1QTY-REGULER (4 pesanan), 1 picklist "
          "KOMBINASI-REGULER (2 pesanan)")

    # bagian="1qty": cuma proses bagian itu
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            hasil = pl.proses_reguler(k, resi_spesial_semua, Path(d) / "r.xlsx",
                                      Path(d) / "label", bagian="1qty")
    finally:
        pl.lanjutkan_picklist = asli
    assert {h["SKU"] for h in hasil} == {"1QTY-REGULER"}, hasil
    print("  bagian bisa dibatasi 1qty atau kombinasi saja (dipakai --bagian)")


class JubelioPalsuShopeePagi:
    """Server tiruan khusus skenario picklist Shopee Pagi (filter channel + jam pesan)."""

    def __init__(self):
        self.log = []
        # source 64 = Shopee (izin); source 4 = Lazada nyelip (harus dibuang lagi di kita)
        self.pesanan = [
            {"salesorder_id": 1, "source": 64, "grand_total": "35900.0000", "transaction_date": "2026-09-29T02:00:00.000Z"},  # 09:00 WIB
            {"salesorder_id": 2, "source": 64, "grand_total": "35900.0000", "transaction_date": "2026-09-29T04:59:59.000Z"},  # 11:59:59 WIB
            {"salesorder_id": 3, "source": 64, "grand_total": "35900.0000", "transaction_date": "2026-09-29T05:00:00.000Z"},  # 12:00:00 WIB, pas batas -> ikut
            {"salesorder_id": 4, "source": 64, "grand_total": "35900.0000", "transaction_date": "2026-09-29T05:00:01.000Z"},  # 12:00:01 WIB -> tidak ikut
            {"salesorder_id": 5, "source": 64, "grand_total": "35900.0000", "transaction_date": "2026-09-29T06:00:00.000Z"},  # 13:00 WIB -> tidak ikut
            {"salesorder_id": 6, "source": 4, "grand_total": "35900.0000", "transaction_date": "2026-09-29T02:00:00.000Z"},   # Lazada nyelip
        ]
        self.picklist_no = 0

    def get(self, url, params=None, headers=None, timeout=None, cookies=None):
        self.log.append(("GET", url, params))
        assert headers.get("authorization") == "TKN"
        if url.endswith("ready-to-process/"):
            assert params.get("channel_ids[0]") == pl.CHANNEL_ID_SHOPEE
            assert "channel_ids[1]" not in params, "Shopee Pagi cuma channel Shopee"
            assert [params[f"order_type[{i}]"] for i in range(len(pl.TIPE_PESANAN_FILTER))] \
                == pl.TIPE_PESANAN_FILTER, "channel Shopee -> pesanan kilat harus difilter"
            return Resp(data={"data": self.pesanan, "totalCount": len(self.pesanan)})
        raise AssertionError(f"GET tak dikenal {url}")

    def post(self, url, json=None, headers=None, timeout=None, cookies=None):
        self.log.append(("POST", url, json))
        assert headers.get("authorization") == "TKN"
        if url.endswith("items-to-pick/"):
            return Resp(data=[{"salesorder_detail_id": 9000 + i, "item_id": 1, "location_id": -1,
                               "qty_ordered": "1.0000", "salesorder_id": i, "bundle_item_id": 0,
                               "package_detail_id": None, "package_id": None, "end_qty": "999.0000",
                               "item_full_name": "X", "salesorder_no": f"SO{i}"} for i in json["ids"]])
        if url.endswith("wms/sales/picklists/"):
            self.picklist_no += 1
            no = f"PICK-{920000 + self.picklist_no}"
            return Resp(data={"status": "ok", "data": {
                "picks": [{"picklist_id": self.picklist_no, "picklist_no": no, "status": "ok"}],
                "invalidSO": []}})
        raise AssertionError(f"POST tak dikenal {url}")


def uji_shopee_pagi_filter_channel_dan_jam_cutoff():
    import datetime as dt

    j = JubelioPalsuShopeePagi()
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    # "sekarang" = jam 13:00 WIB 29-09-2026 -> batas = 12:00 WIB hari yang sama
    sekarang = dt.datetime(2026, 9, 29, 13, 0, 0, tzinfo=pl.WIB)

    pesanan = pl.ambil_pesanan_shopee_pagi(k, sekarang=sekarang)
    assert {o["salesorder_id"] for o in pesanan} == {1, 2, 3}, pesanan
    print("  Shopee Pagi: hanya channel Shopee (Lazada dibuang), hanya jam pesan <= 12:00:00 "
          "WIB (12:00:00 pas ikut, 12:00:01 tidak)")

    panggilan = []
    asli = pl.lanjutkan_picklist

    def stub(k, pid, pno, jumlah, sku, folder_label, nama_file=None):
        panggilan.append((pid, pno, jumlah, sku, folder_label))
        return {"Waktu": "-", "SKU": sku, "No Picklist": pno, "Total Pesanan": jumlah,
                "Resi Keluar": jumlah, "File Label": f"{pno}_{sku}_x.pdf", "Catatan": ""}
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            riwayat = Path(d) / "riwayat.xlsx"
            # proses_shopee_pagi() sendiri pakai waktu sungguhan (dipanggil manual jam 13:00,
            # bukan lewat parameter) -> di sini cukup uji lewat pemanggilan langsung supaya
            # deterministik, ambil_pesanan_shopee_pagi() sudah diuji terpisah di atas
            pesanan = pl.ambil_pesanan_shopee_pagi(k, sekarang=sekarang)
            hasil = pl._proses_channel_batch(k, "Shopee Pagi", pl.LABEL_SHOPEE_PAGI, pesanan,
                                             riwayat, Path(d) / "label")
    finally:
        pl.lanjutkan_picklist = asli

    assert len(hasil) == 1 and hasil[0]["SKU"] == "SHOPEE-PAGI" and hasil[0]["Total Pesanan"] == 3
    assert panggilan[0][3] == "SHOPEE-PAGI"
    print("  proses_shopee_pagi: 3 pesanan -> 1 picklist SHOPEE-PAGI")


class JubelioPalsuJntSiang:
    """Server tiruan khusus skenario picklist J&T Resi Siang (filter channel + kurir +
    jam pesan)."""

    def __init__(self):
        self.log = []
        # channel 131076 = TikTok Shop (izin); channel 64 = Shopee nyelip (harus dibuang lagi)
        self.pesanan = [
            {"salesorder_id": 1, "source": 131076, "grand_total": "35900.0000", "transaction_date": "2026-09-29T02:00:00.000Z"},  # 09:00 WIB
            {"salesorder_id": 2, "source": 131076, "grand_total": "35900.0000", "transaction_date": "2026-09-29T07:59:59.000Z"},  # 14:59:59 WIB
            {"salesorder_id": 3, "source": 131076, "grand_total": "35900.0000", "transaction_date": "2026-09-29T08:00:00.000Z"},  # 15:00:00 WIB, pas batas -> ikut
            {"salesorder_id": 4, "source": 131076, "grand_total": "35900.0000", "transaction_date": "2026-09-29T08:00:01.000Z"},  # 15:00:01 WIB -> tidak ikut
            {"salesorder_id": 5, "source": 131076, "grand_total": "35900.0000", "transaction_date": "2026-09-29T09:00:00.000Z"},  # 16:00 WIB -> tidak ikut
            {"salesorder_id": 6, "source": 64, "grand_total": "35900.0000", "transaction_date": "2026-09-29T02:00:00.000Z"},      # Shopee nyelip
        ]
        self.picklist_no = 0

    def get(self, url, params=None, headers=None, timeout=None, cookies=None):
        self.log.append(("GET", url, params))
        assert headers.get("authorization") == "TKN"
        if url.endswith("ready-to-process/"):
            assert params.get("channel_ids[0]") == pl.CHANNEL_ID_TIKTOK_SHOP
            assert "channel_ids[1]" not in params, "J&T Resi Siang cuma channel TikTok Shop"
            assert params.get("couriers[0]") == "j&t" and "couriers[1]" not in params, \
                "J&T Resi Siang cuma kurir J&T"
            return Resp(data={"data": self.pesanan, "totalCount": len(self.pesanan)})
        raise AssertionError(f"GET tak dikenal {url}")

    def post(self, url, json=None, headers=None, timeout=None, cookies=None):
        self.log.append(("POST", url, json))
        assert headers.get("authorization") == "TKN"
        if url.endswith("items-to-pick/"):
            return Resp(data=[{"salesorder_detail_id": 9000 + i, "item_id": 1, "location_id": -1,
                               "qty_ordered": "1.0000", "salesorder_id": i, "bundle_item_id": 0,
                               "package_detail_id": None, "package_id": None, "end_qty": "999.0000",
                               "item_full_name": "X", "salesorder_no": f"SO{i}"} for i in json["ids"]])
        if url.endswith("wms/sales/picklists/"):
            self.picklist_no += 1
            no = f"PICK-{930000 + self.picklist_no}"
            return Resp(data={"status": "ok", "data": {
                "picks": [{"picklist_id": self.picklist_no, "picklist_no": no, "status": "ok"}],
                "invalidSO": []}})
        raise AssertionError(f"POST tak dikenal {url}")


def uji_jnt_siang_filter_channel_kurir_dan_jam_cutoff():
    import datetime as dt

    j = JubelioPalsuJntSiang()
    k = pl.Klien("TKN", sesi=j, tidur=lambda s: None)
    # "sekarang" = jam 16:00 WIB 29-09-2026 -> batas = 15:00 WIB hari yang sama
    sekarang = dt.datetime(2026, 9, 29, 16, 0, 0, tzinfo=pl.WIB)

    pesanan = pl.ambil_pesanan_jnt_siang(k, sekarang=sekarang)
    assert {o["salesorder_id"] for o in pesanan} == {1, 2, 3}, pesanan
    print("  J&T Resi Siang: hanya channel TikTok Shop + kurir J&T (Shopee dibuang), hanya "
          "jam pesan <= 15:00:00 WIB (15:00:00 pas ikut, 15:00:01 tidak)")

    panggilan = []
    asli = pl.lanjutkan_picklist

    def stub(k, pid, pno, jumlah, sku, folder_label, nama_file=None):
        panggilan.append((pid, pno, jumlah, sku, folder_label))
        return {"Waktu": "-", "SKU": sku, "No Picklist": pno, "Total Pesanan": jumlah,
                "Resi Keluar": jumlah, "File Label": f"{pno}_{sku}_x.pdf", "Catatan": ""}
    pl.lanjutkan_picklist = stub
    try:
        with tempfile.TemporaryDirectory() as d:
            riwayat = Path(d) / "riwayat.xlsx"
            # proses_jnt_siang() sendiri pakai waktu sungguhan (dipanggil manual jam 15:00,
            # bukan lewat parameter) -> di sini cukup uji lewat pemanggilan langsung supaya
            # deterministik, ambil_pesanan_jnt_siang() sudah diuji terpisah di atas
            pesanan = pl.ambil_pesanan_jnt_siang(k, sekarang=sekarang)
            hasil = pl._proses_channel_batch(k, "J&T Resi Siang", pl.LABEL_JNT_SIANG, pesanan,
                                             riwayat, Path(d) / "label")
    finally:
        pl.lanjutkan_picklist = asli

    assert len(hasil) == 1 and hasil[0]["SKU"] == "JNT-SIANG" and hasil[0]["Total Pesanan"] == 3
    assert panggilan[0][3] == "JNT-SIANG"
    print("  proses_jnt_siang: 3 pesanan -> 1 picklist JNT-SIANG")


def uji_cek_item_bundle_hanya_ptaa_boleh_spesial():
    item = lambda so, bundle: {"salesorder_id": so, "salesorder_detail_id": so, "qty_ordered": "1.0000",
                               "bundle_item_id": bundle, "item_full_name": "komponen tidak terkait",
                               "location_id": -1}
    # bundle (bundle_item_id != 0) + SKU mengandung PTAA -> lolos, tidak ada exception
    pl._cek_item("T01-PTAA-5", [1], [item(1, 77)])
    pl._cek_item("t01-ptaa-20", [1], [item(1, 77)])   # tidak peka huruf besar/kecil
    print("  bundle SKU mengandung PTAA: lolos _cek_item()")

    # bundle tapi SKU BUKAN PTAA (mis. PTAE, PTAD, BKAG) -> Lewati, tidak boleh jadi spesial
    for sku in ("T01-PTAE-2", "T01-PTAD-3", "T01-BKAG-2"):
        try:
            pl._cek_item(sku, [1], [item(1, 77)])
        except pl.Lewati:
            pass
        else:
            raise AssertionError(f"{sku}: bundle non-PTAA seharusnya Lewati")
    print("  bundle SKU bukan PTAA (PTAE/PTAD/BKAG): Lewati, tidak boleh jadi spesial")

    # non-bundle (bundle_item_id 0) tidak terpengaruh aturan PTAA sama sekali
    pl._cek_item("C227-BRS3-2", [1], [{"salesorder_id": 1, "salesorder_detail_id": 1,
                                       "qty_ordered": "1.0000", "bundle_item_id": 0,
                                       "item_full_name": "C227-BRS3-2 - Baju", "location_id": -1}])
    print("  SKU non-bundle: aturan PTAA tidak berlaku, tetap lolos seperti biasa")


JEDA_RESI = pl.JEDA_RESI_S

if __name__ == "__main__":
    import logging
    logging.basicConfig(level=logging.WARNING)
    for nama, f in list(globals().items()):
        if nama.startswith("uji_"):
            print(nama)
            f()
    print("SEMUA UJI LULUS")
