"""Uji rekap_master_excel.py: PICKLIST.xlsx per sesi dari data/template/picklist-form-kosong.xlsx.
Tidak menyentuh file master tim sama sekali (modulnya memang tidak pernah membukanya).

Jalankan:  .venv\\Scripts\\python tests\\test_rekap_master_excel.py
"""
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import operator_aktif  # noqa: E402
import rekap_master_excel as rme  # noqa: E402
import peringatan_gagal  # noqa: E402

KUNING = "FFFFFF00"


def _reset():
    rme._buku.clear()
    rme._simpan_pernah_gagal = False
    peringatan_gagal._file = None
    operator_aktif.FILE = Path(tempfile.gettempdir()) / "tidak-ada-operator.json"   # -> bawaan PUTRI


def _baris(no: int, waktu: str = "06-10-2026 13:45", **lain) -> dict:
    return {"Waktu": waktu, "SKU": f"SKU{no}", "No Picklist": f"PICK-000{no}",
            "Total Pesanan": 3, "Resi Keluar": 3, **lain}


def _muat(folder: Path):
    from openpyxl import load_workbook

    wb = load_workbook(folder / rme.NAMA_FILE)
    return wb, wb.worksheets[0]


def _kuning(sel) -> bool:
    return sel.fill.fill_type == "solid" and sel.fill.fgColor.rgb == KUNING


def uji_template_sama_dengan_form_kosong_tim():
    from openpyxl import load_workbook

    wb = load_workbook(rme.TEMPLATE)
    assert wb.sheetnames == ["HARI IN - FORM KOSONG"], wb.sheetnames
    ws = wb.worksheets[0]
    assert ws["E2"].value == "PICK LIST - HARI INI" and ws["M2"].value == "=SUBTOTAL(9,M6:M40000)"
    assert [ws.cell(4, c).value for c in (5, 6, 7, 8, 12, 13, 14, 15, 16, 17)] == [
        "NO", "OPR", "TGL PROSES", "JAM", "NO PICK LIST", "LOLOS", "PRINT", "MINUS", "SCAN", "CTRL"]
    assert ws["O6"].value == '=IF(L6="","",IF(N6="","",M6-N6))'
    assert ws["E6"].value == '=IF(L6="","",IF(G6=G5,E5+1,1))'
    assert ws["Q6"].value == '=IF(P6="","",M6-P6)'
    assert ws["B6"].value == '=IF(L6="","",I6&"-"&J6)' and ws["W6"].value == '=IF(L6="","",G6)'
    assert ws.column_dimensions["B"].hidden and ws.freeze_panes == "A6"
    print("  template: sheet, header, rumus B/E/O/Q/W & subtotal sama dengan form kosong tim")


def uji_catat_buat_file_sesi_dari_template_rumus_dan_format_ikut():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        rme.catat(folder, _baris(157401))
        rme.catat(folder, _baris(157402, **{"Resi Keluar": 3, "Tanpa Resi": ["SO1", "SO2"]}))
        rme.catat(folder, _baris(157403, "06-10-2026 18:53"))
        assert not (folder / rme.NAMA_SEMENTARA).exists()
        wb, ws = _muat(folder)
        assert wb.sheetnames == ["HARI IN - FORM KOSONG"]
        assert [ws.cell(r, 12).value for r in (6, 7, 8, 9)] == [157401, 157402, 157403, None]
        assert ws.cell(6, 6).value == "PUTRI" and ws.cell(6, 7).value == datetime(2026, 10, 6)
        assert ws.cell(6, 8).value == 13.45 and ws.cell(8, 8).value == 18.53
        assert (ws.cell(6, 13).value, ws.cell(6, 14).value, ws.cell(6, 20).value) == (3, 3, "SKU157401")
        assert ws.cell(7, 21).value == "SO1, SO2" and ws.cell(6, 21).value is None
        # rumus disalin ke baris baru (relatif) & tidak ada nilai di kolom rumus
        assert ws["O8"].value == '=IF(L8="","",IF(N8="","",M8-N8))'
        assert ws["E8"].value == '=IF(L8="","",IF(G8=G7,E7+1,1))'
        assert ws["Q7"].value == '=IF(P7="","",M7-P7)' and ws["W8"].value == '=IF(L8="","",G8)'
        assert ws["B7"].value == '=IF(L7="","",I7&"-"&J7)'
        # format: font, number format, tinggi baris sama dengan baris contoh
        assert ws["M8"].number_format == "#,##0" and ws["G8"].number_format == ws["G6"].number_format
        assert ws["L8"].font.name == "Cambria" and ws["L8"].font.b and ws["L8"].border.left.style == "thin"
        assert ws.row_dimensions[8].height == ws.row_dimensions[6].height
        assert not any(_kuning(ws.cell(r, c)) for r in (6, 7, 8) for c in range(2, 22))
    print("  catat(): file sesi dari template, nilai F/G/H/L/M/N/T/U terisi, rumus & format disalin")


def uji_catatan_non_wajib_keluar_kolom_u():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        rme.atur_catatan_proses("NON WAJIB KELUAR")
        try:
            rme.catat(folder, _baris(157601))
            rme.catat(folder, _baris(157602, **{"Tanpa Resi": ["SO1", "SO2"]}), [157603])
        finally:
            rme.atur_catatan_proses("")
        rme.catat(folder, _baris(157604, **{"Tanpa Resi": ["SO9"]}))
        _, ws = _muat(folder)
        assert ws.cell(6, 21).value == "NON WAJIB KELUAR"
        assert ws.cell(7, 12).value == 157603 and ws.cell(7, 21).value is None     # PICKLIST CANCEL tanpa catatan
        assert ws.cell(8, 21).value == "NON WAJIB KELUAR | SO1, SO2"
        assert ws.cell(9, 21).value == "SO9"                                       # setelah dimatikan
    print("  kolom U: NON WAJIB KELUAR (digabung ' | ' dgn tanpa resi), tanpa awalan setelah dimatikan")


def uji_minus_lebih_dari_nol_sel_o_kuning():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        rme.catat(folder, _baris(157501, **{"Total Pesanan": 10, "Resi Keluar": 8}))     # minus 2
        rme.catat(folder, _baris(157502))                                              # minus 0
        _, ws = _muat(folder)
        assert _kuning(ws["O6"]) and not _kuning(ws["O7"])
        assert not any(_kuning(ws.cell(6, c)) for c in range(2, 22) if c != 15)       # hanya sel O
    print("  MINUS > 0: hanya sel O baris itu yang kuning")


def uji_picklist_cancel_hanya_nomor_terlompat_yang_diberikan():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        rme.catat(folder, _baris(157601))
        rme.catat(folder, _baris(157604), [157602, 157603])
        _, ws = _muat(folder)
        assert [ws.cell(r, 12).value for r in (6, 7, 8, 9)] == [157601, 157602, 157603, 157604]
        for r in (7, 8):
            assert ws.cell(r, 20).value == "PICKLIST CANCEL"
            assert ws.cell(r, 6).value == "PUTRI" and ws.cell(r, 7).value == datetime(2026, 10, 6)
            assert all(_kuning(ws.cell(r, c)) for c in range(2, 22)), r
            assert not _kuning(ws.cell(r, 1))
        assert not any(_kuning(ws.cell(9, c)) for c in range(2, 22))
        assert ws["E9"].value == '=IF(L9="","",IF(G9=G8,E8+1,1))'
    print("  nomor_terlompat: baris kuning PICKLIST CANCEL sebelum baris picklist, rumus tetap")


def uji_catatan_gagal_atau_terhenti_baris_kuning():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        rme.catat(folder, {"Waktu": "06-10-2026 14:00", "SKU": "X", "No Picklist": "PICK-000157701",
                           "Total Pesanan": 5, "Catatan": "TERHENTI: timeout. Lanjutkan: ..."})
        _, ws = _muat(folder)
        assert all(_kuning(ws.cell(6, c)) for c in range(2, 22))
        assert ws.cell(6, 14).value is None
    print("  Catatan GAGAL/TERHENTI: seluruh baris picklist kuning")


def uji_tanpa_no_picklist_dilewati_dan_tanpa_file():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        rme.catat(folder, {"Waktu": "06-10-2026 14:00", "SKU": "X", "Catatan": "Dilewati: tidak ada"})
        assert not (folder / rme.NAMA_FILE).exists()
    print("  catat(): baris tanpa No Picklist (gagal sebelum picklist dibuat) dilewati")


def uji_sesi_berbeda_file_berbeda_dan_lanjut_di_file_yang_sudah_ada():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        a, b = Path(tmp) / "1", Path(tmp) / "2"
        a.mkdir(), b.mkdir()
        rme.catat(a, _baris(157801))
        rme.catat(b, _baris(157802))
        rme.catat(a, _baris(157803))
        assert [_muat(a)[1].cell(r, 12).value for r in (6, 7)] == [157801, 157803]
        assert _muat(b)[1].cell(6, 12).value == 157802
        # proses python baru (langkah .bat berikutnya) membuka file sesi yang sama dan melanjutkan
        _reset()
        rme.catat(a, _baris(157804), [])
        assert [_muat(a)[1].cell(r, 12).value for r in (6, 7, 8)] == [157801, 157803, 157804]
        assert _muat(a)[1]["E8"].value == '=IF(L8="","",IF(G8=G7,E7+1,1))'
    print("  1 file per sesi; proses berikutnya melanjutkan di baris setelah data terakhir")


def uji_template_hilang_tidak_menggagalkan_proses():
    _reset()
    asli = rme.TEMPLATE
    try:
        rme.TEMPLATE = Path("tidak/ada/template.xlsx")
        with tempfile.TemporaryDirectory() as tmp:
            rme.catat(Path(tmp), _baris(157901))
            assert not (Path(tmp) / rme.NAMA_FILE).exists()
    finally:
        rme.TEMPLATE = asli
        rme._sudah_peringatan_template = False
    print("  template tidak ada: catat() tidak error, tidak ada file")


def uji_simpan_gagal_file_lama_utuh_dan_tersimpan_di_catat_berikutnya():
    """PICKLIST.xlsx dibuka di Excel => os.replace gagal: file lama utuh, sementara dibuang,
    isi tetap di memori & ikut tersimpan begitu bisa menulis lagi."""
    _reset()
    asli = rme._ganti_file
    peringatan = []
    asli_catat = peringatan_gagal.catat
    peringatan_gagal.catat = peringatan.append
    try:
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            rme.catat(folder, _baris(158001))

            def terkunci(asal, tujuan, coba_maks):
                raise PermissionError("dibuka di Excel")
            rme._ganti_file = terkunci
            rme.catat(folder, _baris(158002))
            assert not (folder / rme.NAMA_SEMENTARA).exists()
            assert [_muat(folder)[1].cell(r, 12).value for r in (6, 7)] == [158001, None]   # lama utuh
            rme.selesai()                                   # masih terkunci -> peringatan
            assert len(peringatan) == 1 and "belum tersimpan" in peringatan[0], peringatan
            rme._ganti_file = asli                          # Excel ditutup
            rme.catat(folder, _baris(158003))
            assert [_muat(folder)[1].cell(r, 12).value for r in (6, 7, 8)] == [158001, 158002, 158003]
            rme.selesai()
            assert len(peringatan) == 1
    finally:
        rme._ganti_file = asli
        peringatan_gagal.catat = asli_catat
    print("  simpan gagal: file lama utuh, isi ikut tersimpan lagi di catat() berikutnya, "
          "peringatan kalau tetap gagal di akhir proses")


def uji_file_rusak_dipindah_lalu_dibuat_baru():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        (folder / rme.NAMA_FILE).write_bytes(b"bukan zip")
        rme.catat(folder, _baris(158101))
        assert _muat(folder)[1].cell(6, 12).value == 158101
        assert list(folder.glob("PICKLIST_rusak_*.xlsx"))
    print("  PICKLIST.xlsx rusak: dipindah ke PICKLIST_rusak_*, dibuat baru dari template")


def uji_cepat_puluhan_picklist():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        mulai = time.monotonic()
        for i in range(60):
            rme.catat(folder, _baris(160000 + i))
        detik = time.monotonic() - mulai
        assert _muat(folder)[1].cell(65, 12).value == 160059
    assert detik < 30, detik
    print(f"  60 picklist ditulis langsung (simpan tiap picklist) dalam {detik:.1f} detik")


def uji_isi_scan_dari_iresis():
    _reset()
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        assert rme.isi_scan(folder, {157401: 5}) == 0 and not (folder / rme.NAMA_FILE).exists()
        rme.catat(folder, _baris(157401))
        rme.catat(folder, _baris(157403), [157402])
        assert rme.isi_scan(folder, {157401: 5, 157403: 7, 999: 1}) == 2
        ws = _muat(folder)[1]
        assert [ws.cell(r, 16).value for r in (6, 7, 8)] == [5, None, 7]     # P=SCAN; cancel kosong
        assert ws["Q6"].value == '=IF(P6="","",M6-P6)'
        assert rme.isi_scan(folder, {157401: 5, 157403: 7}) == 0              # tak berubah
        assert rme.isi_scan(folder, {157401: 6}) == 1 and _muat(folder)[1]["P6"].value == 6
    print("  isi_scan(): kolom P diisi dari total resi IRESIS per nomor picklist, idempoten")


if __name__ == "__main__":
    for nama, f in list(globals().items()):
        if nama.startswith("uji_"):
            print(nama)
            f()
    print("SEMUA UJI LULUS")
