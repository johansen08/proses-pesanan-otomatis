"""Uji pilah_satuan.py tanpa API/Jubelio sungguhan (klien tiruan, PDF label tiruan).

Jalankan:  .venv\\Scripts\\python tests\\test_pilah_satuan.py
"""
import json
import sys
import tempfile
from pathlib import Path

from pypdf import PdfReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import kurasi_satuan  # noqa: E402
import proses_label  # noqa: E402
import pilah_satuan as ps  # noqa: E402
import print_spesial  # noqa: E402

# Tes TIDAK boleh menulis catatan sungguhan di logs/: arahkan ke folder sementara.
_LOG_TES = Path(tempfile.mkdtemp())
kurasi_satuan.FILE_KURASI = _LOG_TES / "satuan_dikurasi.json"
print_spesial.FILE_SUDAH_DICETAK = _LOG_TES / "sudah_dicetak.txt"
ps.MASA_TENANG_S = 0   # file tes baru dibuat; cek "sesi masih menulis" diuji khusus di bawah


def _pdf_label(tujuan: Path, halaman: list[tuple[str, str]]) -> None:
    c = canvas.Canvas(str(tujuan))
    for pesanan, resi in halaman:
        c.drawString(50, 700, "No Pesanan")
        c.drawString(50, 680, f": {pesanan}")
        c.drawString(50, 660, f"No Resi: {resi}")
        c.showPage()
    c.save()


class KlienPalsu:
    """item per picklist: {nomor: [(sku, pesanan, resi, status), ...]}"""

    def __init__(self, data):
        self.data = data

    def get(self, path):
        nomor = int(path.rsplit("/", 1)[1])
        return {"items": [{"item_code": s, "salesorder_no": o, "tracking_no": r,
                           "shipper": "SPX Hemat", "wms_status": w}
                          for s, o, r, w in self.data[nomor]]}


def _siapkan(tmp: Path):
    # sesi 1: picklist 10 (X x2, Y x1, Z batal); sesi 2: picklist 20 (X x1, Y x1) -> X=3, Y=2
    (tmp / "1" / "SATUAN").mkdir(parents=True)
    (tmp / "2" / "SATUAN").mkdir(parents=True)
    _pdf_label(tmp / "1" / "SATUAN" / "PICK-000000010_1QTY-REGULER-2A_2026-10-13_081500.pdf",
               [("O1", "R1"), ("O2", "R2"), ("O3", "R3"), ("O4", "R4")])
    _pdf_label(tmp / "2" / "SATUAN" / "PICK-000000020_1QTY-REGULER-2A_2026-10-13_104000.pdf",
               [("O5", "R5"), ("O6", "R6")])
    return KlienPalsu({
        10: [("X", "O1", "R1", "FINISH_PICK"), ("X", "O2", "R2", "FINISH_PICK"),
             ("Y", "O3", "R3", "FINISH_PICK"), ("Z", "O4", "R4", "CANCELED")],
        20: [("X", "O5", "R5", "FINISH_PICK"), ("Y", "O6", "R6", "FINISH_PICK")]})


def uji_semua_jenis_satuan_gabungan_dan_per_kurir_ikut():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        _siapkan(tmp)
        (tmp / "1" / "JNT_SATUAN").mkdir()
        (tmp / "1" / "SPXHEMAT_SATUAN").mkdir()
        _pdf_label(tmp / "1" / "JNT_SATUAN" / "PICK-000000030_1QTY-REGULER-2A_2026-10-13_082000.pdf", [("O7", "R7")])
        _pdf_label(tmp / "1" / "SPXHEMAT_SATUAN" / "PICK-000000031_1QTY-REGULER-2A_2026-10-13_082100.pdf", [("O8", "R8")])
        assert [f.name[:14] for f in ps.cari_pdf_satuan(tmp)] == [
            "PICK-000000010", "PICK-000000020", "PICK-000000030", "PICK-000000031"]
        assert [f.name[:14] for f in ps.cari_pdf_satuan(tmp, ["1"])] == [
            "PICK-000000010", "PICK-000000030", "PICK-000000031"]
        # sesi yang HANYA punya satuan per kurir: boleh dipilah, jam berapa pun (tidak kena syarat jam)
        (tmp / "3" / "JNT_SATUAN").mkdir(parents=True)
        _pdf_label(tmp / "3" / "JNT_SATUAN" / "PICK-000000040_JNT-1QTY-REGULER-2A_2026-10-13_153000.pdf", [("O9", "R9")])
        assert ps.periksa_sesi(tmp / "3") == (True, None)
        assert ps.pdf_terpilih(tmp, ["1", "3"])[1] == ["1", "3"]
        assert ps.pdf_terpilih(tmp)[1] == ["1", "2", "3"]
        # sesi gabungan sore hari juga boleh (tidak ada batas jam)
        (tmp / "4" / "SATUAN").mkdir(parents=True)
        _pdf_label(tmp / "4" / "SATUAN" / "PICK-000000050_1QTY-REGULER-2A_2026-10-13_153000.pdf", [("O10", "R10")])
        assert ps.periksa_sesi(tmp / "4") == (True, None)
        assert ps.pdf_terpilih(tmp)[1] == ["1", "2", "3", "4"]
        # kelompok rak satuan per kurir dinormalkan ke kelompok yang sama dengan gabungan
        assert ps.nama_kelompok(Path("PICK-000000040_JNT-1QTY-REGULER-2A_2026-10-13_153000.pdf")) == "1QTY-REGULER-2A"
        assert ps.nama_kelompok(Path("PICK-000000040_SPXHEMAT-1QTY-REGULER-LAINNYA_2026-10-13_153000.pdf")) == "1QTY-REGULER-LAINNYA"
        assert ps.nama_kelompok(Path("PICK-000000010_1QTY-REGULER-3B_2026-10-13_081500.pdf")) == "1QTY-REGULER-3B"
    print("  semua jenis satuan ikut (SATUAN, JNT_/SPXHEMAT_SATUAN); semua sesi bebas jam")


def uji_satuan_jnt_siang_dipilah_bersama_satuan_jam_sebelumnya():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        for sub in ("1/SATUAN", "4/SATUAN", "4/JNTSIANG_SATUAN", "5/SATUAN"):
            (tmp / sub).mkdir(parents=True)
        # sesi 1 (pagi 08.15, gabungan) + sesi 4 (15.00: gabungan di luar jam + JNTSIANG_SATUAN per kurir)
        _pdf_label(tmp / "1" / "SATUAN" / "PICK-000000010_1QTY-REGULER-2A_2026-10-13_081500.pdf", [("O1", "R1"), ("O2", "R2")])
        _pdf_label(tmp / "4" / "SATUAN" / "PICK-000000040_1QTY-REGULER-3A_2026-10-13_150000.pdf", [("O3", "R3")])
        _pdf_label(tmp / "4" / "JNTSIANG_SATUAN" / "PICK-000000041_JNTSIANG-1QTY-REGULER-2A_2026-10-13_150100.pdf",
                   [("O4", "R4"), ("O5", "R5")])
        # tidak ada batas jam: sesi 4 (15.00) boleh, gabungan maupun per kurir
        assert ps.periksa_sesi(tmp / "4") == (True, None)
        pdfs, sesi = ps.pdf_terpilih(tmp)
        assert sesi == ["1", "4"]
        nama = [f.name[:14] for f in pdfs]
        assert nama == ["PICK-000000010", "PICK-000000040", "PICK-000000041"], nama   # gabungan sesi 4 ikut
        assert ps.nama_kelompok(pdfs[2]) == "1QTY-REGULER-2A"
        assert len(ps.cari_pdf_satuan(tmp)) == 3
        # sesi yang HANYA punya satuan gabungan sore hari juga boleh
        _pdf_label(tmp / "5" / "SATUAN" / "PICK-000000050_1QTY-REGULER-2A_2026-10-13_160000.pdf", [("O6", "R6")])
        assert ps.periksa_sesi(tmp / "5") == (True, None)
    print("  satuan J&T siang (JNTSIANG_SATUAN, sesi 15.00) dipilah bersama satuan jam sebelumnya; gabungan sesi itu juga ikut")


def uji_sesi_dari_jam_berapa_pun_boleh_dipilah():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)

        def sesi(no, jam):
            f = tmp / no / "SATUAN"
            f.mkdir(parents=True)
            _pdf_label(f / f"PICK-0000000{no}0_1QTY-REGULER-2A_2026-10-13_{jam}.pdf", [("O" + no, "R" + no)])
        kasus = {"1": "070000", "2": "120100", "3": "153000", "4": "195900", "5": "200000", "6": "003000"}
        for no, jam in kasus.items():
            sesi(no, jam)
        assert all(ps.periksa_sesi(tmp / no) == (True, None) for no in kasus)
        # tanpa cap waktu di nama file pun boleh
        (tmp / "7" / "SATUAN").mkdir(parents=True)
        _pdf_label(tmp / "7" / "SATUAN" / "PICK-000000070_1QTY-REGULER-2A_x.pdf", [("O8", "R8")])
        assert ps.periksa_sesi(tmp / "7") == (True, None)
        assert ps.pdf_terpilih(tmp, ["1", "4"])[1] == ["1", "4"]
        assert ps.pdf_terpilih(tmp)[1] == ["1", "2", "3", "4", "5", "6", "7"]
        # sesi tanpa label satuan tetap ditolak
        (tmp / "8").mkdir()
        ok, alasan = ps.periksa_sesi(tmp / "8")
        assert not ok and "satuan" in alasan
    print("  sesi dari jam berapa pun (juga tanpa cap waktu) boleh dipilah; tanpa label satuan ditolak")


def uji_sesi_pemulihan_dipakai_pilah_dengan_jam_sesi_asal():
    assert ps.parse_sesi("1-2-11P") == ["1", "2", "11P"] and ps.parse_sesi("11-PEMULIHAN") == ["11P"]
    assert ps.parse_sesi("11p-3") == ["3", "11P"]
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)

        def sesi(nama, jam, picklist, tanggal="2026-10-19"):
            f = tmp / nama / "SATUAN"
            f.mkdir(parents=True)
            _pdf_label(f / f"PICK-{picklist:09d}_1QTY-REGULER-2A_{tanggal}_{jam}.pdf", [("O" + nama, "R" + nama)])
        sesi("10", "170000", 1)                 # sore
        sesi("11", "072000", 2)                 # pagi: boleh
        sesi("11-PEMULIHAN", "180000", 2)       # dibuat ulang sore hari
        (tmp / "11-PEMULIHAN" / "pulihan.json").write_text(json.dumps({"asal": "2026-10-19/11"}), encoding="utf-8")
        # penanda menunjuk folder <tanggal>/<asal>: letakkan sesi di folder tanggal sungguhan
        tgl = tmp.parent / "2026-10-19"
        assert ps.sesi_asal(tmp / "11-PEMULIHAN") == tmp / "11-PEMULIHAN", "asal tidak ada di bawah folder tanggal -> diri sendiri"
    with tempfile.TemporaryDirectory() as d:
        tgl = Path(d) / "2026-10-19"
        for nama, jam, pk in (("10", "170000", 1), ("11", "072000", 2), ("11-PEMULIHAN", "180000", 2)):
            f = tgl / nama / "SATUAN"
            f.mkdir(parents=True)
            _pdf_label(f / f"PICK-{pk:09d}_1QTY-REGULER-2A_2026-10-19_{jam}.pdf", [("O" + nama, "R" + nama)])
        (tgl / "11-PEMULIHAN" / "pulihan.json").write_text(json.dumps({"asal": "2026-10-19/11"}), encoding="utf-8")
        assert ps.sesi_asal(tgl / "11-PEMULIHAN") == tgl / "11"
        assert ps.periksa_sesi(tgl / "11-PEMULIHAN")[0]
        assert ps.periksa_sesi(tgl / "10")[0], "sesi sore juga boleh"
        pdfs, kode = ps.pdf_terpilih(tgl, ["11P"])
        assert kode == ["11P"] and [f.parent.parent.name for f in pdfs] == ["11-PEMULIHAN"]
        assert ps.nama_sesi(pdfs) == "11P"
        try:
            ps.pdf_terpilih(tgl, ["11", "11P"])
            assert False, "sesi dan pemulihannya tidak boleh dipilih bersamaan"
        except ps.SesiTidakValid as e:
            assert "pilih salah satu" in str(e), e
        assert ps.pdf_terpilih(tgl)[1] == ["10", "11P"], "tanpa --sesi: pemulihan menggantikan sesi asalnya"
    print("  pilah: sesi N-PEMULIHAN (kode 11P) bisa dipilah, tidak bersamaan dengan sesi asalnya")


def uji_resi_kurir_lain_tidak_dihitung_spesial():
    baris = [{"sku": "X", "resi": f"R{i}", "pesanan": f"O{i}", "kurir": "J&T Express Standard" if i < 2 else "SPX Hemat"}
             for i in range(2)] + [{"sku": "X", "resi": "R9", "pesanan": "O9", "kurir": "SiCepat Reguler"}]
    layak, sisa = ps.kelompokkan(baris, 3)
    assert not layak and sisa == {"X": 2}, (layak, sisa)       # kurir lain tidak ikut menggenapkan jadi 3
    baris.append({"sku": "X", "resi": "R10", "pesanan": "O10", "kurir": "spx standard"})
    layak, _ = ps.kelompokkan(baris, 3)
    assert [b["resi"] for b in layak["X"]] == ["R0", "R1", "R10"]   # huruf kecil tetap cocok
    assert ps.kurir_diizinkan("J&T Express NEXT-DAY DELIVERY") and not ps.kurir_diizinkan("GoTo Logistics GTL") and not ps.kurir_diizinkan(None)
    print("  resi berkurir selain J&T/SPX tidak dihitung spesial (tetap sisa)")


def uji_cari_pdf_urut_nomor_pick():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        _siapkan(tmp)
        (tmp / "1" / "SPESIAL").mkdir()
        _pdf_label(tmp / "1" / "SPESIAL" / "PICK-000000001_SPESIAL_X_x.pdf", [("O9", "R9")])
        nama = [f.name[:14] for f in ps.cari_pdf_satuan(tmp)]
        assert nama == ["PICK-000000010", "PICK-000000020"], nama  # SPESIAL tidak ikut
    print("  hanya subfolder *SATUAN, urut nomor PICK")


def uji_pilah_gabung_lintas_sesi():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        h = ps.pilah(k, tmp)
        assert h["pdf"] == 2 and h["halaman"] == 5 and h["tak_cocok"] == 1, h  # O4 batal
        assert set(h["layak"]) == {"X"} and len(h["layak"]["X"]) == 3, h["layak"]
        assert h["sisa_sku"] == 1  # Y baru 2 resi
        pdf_x = tmp / "SPESIAL-SATUAN-1-2" / "SPESIAL-SATUAN_2A_X_3resi.pdf"
        assert pdf_x.exists() and len(PdfReader(str(pdf_x)).pages) == 3
        teks = "".join(p.extract_text() for p in PdfReader(str(pdf_x)).pages)
        assert all(r in teks for r in ("R1", "R2", "R5")) and "R3" not in teks
        assert (tmp / "SPESIAL-SATUAN-1-2" / "detail-resi-spesial-satuan-1-2.xlsx").exists()
        assert not (tmp / "SPESIAL-SATUAN-1-2" / "SPESIAL-SATUAN_2A_Y_2resi.pdf").exists()
    print("  X gabungan 2 sesi (3 resi) dikelompokkan, Y (2) tidak, item batal dibuang")


def uji_sisa_satuan_per_kelompok_rak():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        h = ps.pilah(k, tmp)
        # sisa = Y (O3 di picklist 10, O6 di picklist 20); Z batal dibuang; X pindah ke spesial
        assert {kel: [b["resi"] for b in bs] for kel, bs in h["sisa"].items()} ==             {"1QTY-REGULER-2A": ["R3", "R6"]}, h["sisa"]
        pdf = tmp / "SISA-SATUAN-1-2" / "SISA-SATUAN_1QTY-REGULER-2A_2resi.pdf"
        assert pdf.exists()
        teks = "".join(p.extract_text() for p in PdfReader(str(pdf)).pages)
        assert "R3" in teks and "R6" in teks and "R1" not in teks
    print("  SISA-SATUAN-<sesi>: halaman non-spesial per kelompok rak, tanpa yang batal")


def uji_pilih_sesi():
    assert ps.parse_sesi("1-2-3") == ["1", "2", "3"] and ps.parse_sesi("3, 1") == ["1", "3"]
    assert ps.parse_sesi("") is None and ps.parse_sesi(None) is None
    try:
        ps.parse_sesi("a-1")
        assert False
    except ValueError:
        pass
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        assert [f.name[:14] for f in ps.cari_pdf_satuan(tmp, ["2"])] == ["PICK-000000020"]
        h = ps.pilah(k, tmp, min_resi=1, buat_pdf=False, sesi=["2"])
        assert h["sesi"] == "2" and h["folder"].name == "SPESIAL-SATUAN-2" and h["halaman"] == 2
        from openpyxl import load_workbook
        ws = load_workbook(h["folder"] / "detail-resi-spesial-satuan-2.xlsx")["Detail"]
        baris = list(ws.iter_rows(values_only=True))
        assert baris[0] == ("No Picklist", "SKU", "No Pesanan", "No Resi")
        assert baris[1] == ("PICK-000000020", "X", "O5", "R5"), baris[1]
        # format SAMA dengan detail-resi-spesial.xlsx (proses_label): 1 sheet "Detail", 4 kolom, tanpa tambahan
        wb = load_workbook(h["folder"] / "detail-resi-spesial-satuan-2.xlsx")
        assert wb.sheetnames == ["Detail"] and ws.max_column == 4, (wb.sheetnames, ws.max_column)
        assert list(baris[0]) == proses_label.KOLOM_DETAIL_SPESIAL
        assert [ws.column_dimensions[c].width for c in "ABCD"] == [18, 16, 22, 18]
        assert all(len(b) == 4 for b in baris)
    print("  --sesi memilih sesi; folder SPESIAL-SATUAN-<sesi>; Excel detail format detail-resi-spesial")


def uji_satuan_asal_diberi_penanda_dikurasi():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        asal = ps.cari_pdf_satuan(tmp)
        ps.pilah(k, tmp, tulis=False)
        ps.pilah(k, tmp, buat_pdf=False)
        assert kurasi_satuan.baca() == {}, "analisis-saja / --tanpa-pdf tidak menandai"
        h = ps.pilah(k, tmp)
        assert h["dikurasi"] == 2
        assert kurasi_satuan.baca() == {str(f.resolve()): "1-2" for f in asal}
        # penanda hanya berlaku selama folder hasilnya ada
        import shutil
        shutil.rmtree(tmp / "SPESIAL-SATUAN-1-2", ignore_errors=True)
        assert kurasi_satuan.baca(), "folder SISA-SATUAN masih ada -> penanda masih berlaku"
        shutil.rmtree(tmp / "SISA-SATUAN-1-2")
        assert kurasi_satuan.baca() == {}
    print("  PDF satuan asal ditandai dikurasi hanya saat folder hasil dibuat; hilang kalau folder hasil dihapus")


def uji_status_tercetak_hasil_pilah_dijaga_saat_dibuat_ulang():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        ps.pilah(k, tmp)
        spesial = tmp / "SPESIAL-SATUAN-1-2" / "SPESIAL-SATUAN_2A_X_3resi.pdf"
        sisa = tmp / "SISA-SATUAN-1-2" / "SISA-SATUAN_1QTY-REGULER-2A_2resi.pdf"
        for f in (spesial, sisa):
            print_spesial.catat_sudah_dicetak(f)
        ps.pilah(k, tmp)   # sama persis -> status tercetak tetap
        assert print_spesial.baca_sudah_dicetak() >= {str(spesial.resolve()), str(sisa.resolve())}
        ps.pilah(k, tmp, min_resi=2)   # X & Y spesial: SISA hilang, SPESIAL X isinya sama
        sudah = print_spesial.baca_sudah_dicetak()
        assert str(sisa.resolve()) not in sudah, "file hilang/berubah: status tercetak harus dilupakan"
        assert str(spesial.resolve()) in sudah, "isi identik: status tercetak tetap"
    print("  dibuat ulang: status tercetak tetap bila isi identik, dilupakan bila file berubah/hilang")


def uji_analisis_saja_cache_dan_ringkasan_json():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        cache = tmp / "cache" / "analisis.json"
        h = ps.pilah(k, tmp, tulis=False, cache=cache)
        assert not (tmp / "SPESIAL-SATUAN-1-2").exists() and not (tmp / "SISA-SATUAN-1-2").exists()  # tidak menulis
        assert cache.exists()
        # tahap "buat": cache dipakai, Jubelio (klien) tidak disentuh sama sekali
        h2 = ps.pilah(None, tmp, cache=cache)
        assert set(h2["layak"]) == {"X"} and h2["tak_cocok"] == 1 and (tmp / "SPESIAL-SATUAN-1-2").is_dir()
        # ambang berubah: tetap dari cache, hasil beda, file lama dibersihkan
        h3 = ps.pilah(None, tmp, min_resi=2, cache=cache)
        assert set(h3["layak"]) == {"X", "Y"}
        h4 = ps.pilah(None, tmp, min_resi=3, cache=cache)
        assert not (tmp / "SPESIAL-SATUAN-1-2" / "SPESIAL-SATUAN_2A_Y_2resi.pdf").exists()   # sisa dari ambang 2
        assert sorted(p.name for p in (tmp / "SPESIAL-SATUAN-1-2").iterdir()) == [
            "SPESIAL-SATUAN_2A_X_3resi.pdf", "detail-resi-spesial-satuan-1-2.xlsx"]
        # PDF berubah -> cache tidak berlaku (None -> baca_cache gagal -> butuh klien)
        pdf = ps.cari_pdf_satuan(tmp)[0]
        pdf.write_bytes(pdf.read_bytes() + b"\n%x")
        assert ps.baca_cache(cache, tmp, ps.cari_pdf_satuan(tmp)) is None
        j = ps.ringkasan_json(h4)
        assert j["spesial"] == [{"sku": "X", "n": 3, "rak": "2A", "sesi": {"1": 2, "2": 1}}], j
        assert j["sisa"] == {"2A": 2} and j["tak_cocok"] == 1 and j["folder"] == [
            "SPESIAL-SATUAN-1-2", "SISA-SATUAN-1-2"]
        json_utuh = ps.PENANDA_JSON + __import__("json").dumps(j)
        assert json_utuh.startswith("##HASIL## ")
    print("  analisis-saja tidak menulis; cache menghindari Jubelio; pilah ulang membersihkan file lama; ringkasan JSON")


def uji_ambang_dan_tanpa_pdf():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        h = ps.pilah(k, tmp, min_resi=2, buat_pdf=False)
        assert set(h["layak"]) == {"X", "Y"}
        assert [f.suffix for f in h["file"]] == [".xlsx"]  # tanpa PDF
    print("  --min-resi 2 memasukkan Y; --tanpa-pdf hanya Excel")


def _baris(rak, resi="R"):
    return {"sku": "?", "resi": resi, "pesanan": resi, "kurir": "SPX Hemat",
            "pdf": Path(f"x/1/SATUAN/PICK-000000001_1QTY-REGULER-{rak}_2026-10-13_081500.pdf")}


def uji_spesial_satuan_urut_rak_lalu_sku():
    layak = {"B": [_baris("2A")] * 3, "A": [_baris("3A")] * 3, "C": [_baris("1B")] * 3, "D": [_baris("2A")] * 3,
             "E": [_baris("2A"), _baris("2A"), _baris("3B")]}     # rak dominan E = 2A
    assert [s for s, _ in ps.urutan_spesial(layak)] == ["C", "B", "D", "E", "A"], ps.urutan_spesial(layak)
    assert ps.rak_dominan(layak["E"]) == "2A"
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        h = ps.pilah(k, tmp)
        assert [f.name for f in h["file"] if f.suffix == ".pdf" and "SPESIAL" in f.name] == ["SPESIAL-SATUAN_2A_X_3resi.pdf"]
        j = ps.ringkasan_json(h)
        assert [(x["rak"], x["sku"]) for x in j["spesial"]] == [("2A", "X")]
    print("  SPESIAL-SATUAN urut rak (1B,2A,2B,3A,3B,LAINNYA) lalu SKU; rak dominan ada di nama file & ringkasan")


def uji_pilah_ulang_di_atas_pdf_yang_sudah_dikurasi_ditolak():
    import shutil
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        ps.pilah(k, tmp, min_resi=1)                      # sesi 1-2 dikurasi -> SPESIAL/SISA-SATUAN-1-2
        asal = ps.cari_pdf_satuan(tmp)
        assert kurasi_satuan.bentrok(asal, "1-2") == {}  # sesi yang persis sama: boleh dibuat ulang
        assert list(kurasi_satuan.bentrok(asal, "1")) == ["1-2"]
        h = ps.pilah(k, tmp, sesi=["1"], tulis=False)     # analisis tetap jalan, hanya diberi peringatan
        assert {k: len(v) for k, v in h["bentrok"].items()} == {"1-2": 1} and ps.ringkasan_json(h)["bentrok"] == {"1-2": 1}
        try:
            ps.pilah(k, tmp, sesi=["1"])
            assert False, "harus ditolak: resi akan ada di dua folder hasil"
        except ps.KurasiBentrok as e:
            assert "SPESIAL-SATUAN-1-2" in str(e) and "hapus folder" in str(e), e
        assert not (tmp / "SPESIAL-SATUAN-1").exists() and not (tmp / "SISA-SATUAN-1").exists()   # tidak ada yang ditulis
        ps.pilah(k, tmp, min_resi=1)                      # sesi sama: tidak bentrok
        ps.pilah(k, tmp, sesi=["1"], buat_pdf=False)      # mode hanya-Excel tidak menandai -> tidak bentrok
        shutil.rmtree(tmp / "SPESIAL-SATUAN-1-2", ignore_errors=True)
        shutil.rmtree(tmp / "SISA-SATUAN-1-2", ignore_errors=True)  # folder lama dihapus -> penanda tak berlaku -> boleh
        assert kurasi_satuan.bentrok(asal, "1") == {}
        ps.pilah(k, tmp, sesi=["1"])
    print("  PDF yang sudah dikurasi ke folder hasil lain: analisis memperingatkan, buat ditolak; sesi sama/folder lama dihapus boleh")


def uji_sesi_yang_masih_menulis_label_ditolak():
    import os
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        _siapkan(tmp)
        lama, ps.MASA_TENANG_S = ps.MASA_TENANG_S, 120
        try:
            ok, alasan = ps.periksa_sesi(tmp / "1")
            assert not ok and "masih menulis label" in alasan, alasan       # file baru dibuat sekarang
            try:
                ps.pdf_terpilih(tmp, ["1"])
                assert False
            except ps.SesiTidakValid as e:
                assert "Sesi 1" in str(e)
            for f in (tmp / "1").rglob("*.pdf"):
                os.utime(f, (f.stat().st_atime, f.stat().st_mtime - 300))   # 5 menit lalu
            assert ps.periksa_sesi(tmp / "1")[0] is True
        finally:
            ps.MASA_TENANG_S = lama
    print("  sesi yang PDF-nya baru diubah (<2 menit) ditolak")


def uji_cache_analisis_kedaluwarsa():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        cache = tmp / "cache.json"
        ps.pilah(k, tmp, tulis=False, cache=cache)
        pdfs = ps.cari_pdf_satuan(tmp)
        assert ps.baca_cache(cache, tmp, pdfs) is not None
        data = json.loads(cache.read_text(encoding="utf-8"))
        data["dibuat"] -= ps.UMUR_CACHE_MAKS_S + 5      # lewat 15 menit
        cache.write_text(json.dumps(data), encoding="utf-8")
        assert ps.baca_cache(cache, tmp, pdfs) is None, "cache kedaluwarsa tidak boleh dipakai"
        del data["dibuat"]                              # cache lama tanpa cap waktu juga ditolak
        cache.write_text(json.dumps(data), encoding="utf-8")
        assert ps.baca_cache(cache, tmp, pdfs) is None
        h = ps.pilah(k, tmp, tulis=False, cache=cache)  # dibaca ulang dari Jubelio & cache diperbarui
        assert ps.baca_cache(cache, tmp, pdfs) is not None and h["tak_cocok"] == 1
    print("  cache analisis kedaluwarsa setelah 15 menit / tanpa cap waktu -> dibaca ulang dari Jubelio")


def uji_nomor_pick_terlompat_di_antara_picklist_satuan():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        for sesi in ("1", "2"):
            (tmp / sesi / "SATUAN").mkdir(parents=True)
        for sesi, nomor in (("1", 10), ("1", 12), ("2", 20), ("2", 21)):
            _pdf_label(tmp / sesi / "SATUAN" / f"PICK-0000000{nomor}_1QTY-REGULER-2A_2026-10-13_081500.pdf", [("O", "R")])
        assert ps.nomor_terlompat_satuan(ps.cari_pdf_satuan(tmp)) == [11]            # 11 hilang di sesi 1
        (tmp / "1" / "KOMBINASI").mkdir()
        _pdf_label(tmp / "1" / "KOMBINASI" / "PICK-000000011_KOMBINASI-REGULER-LANTAI1_2026-10-13_081500.pdf", [("O", "R")])
        assert ps.nomor_terlompat_satuan(ps.cari_pdf_satuan(tmp)) == []              # nomor 11 ada di jenis lain
        h = ps.ringkasan_json({"sesi": "1-2", "pdf": 4, "halaman": 0, "tak_cocok": 0, "kurir_lain": 0, "layak": {},
                               "sisa": {}, "file": [], "terlompat": [11], "bentrok": {}})
        assert h["terlompat"] == ["PICK-000000011"]
    print("  nomor PICK terlompat di antara picklist satuan satu sesi diperingatkan (nomor jenis lain tidak dihitung)")


def uji_mode_hanya_excel_tanpa_pdf_dan_tanpa_penanda():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        h = ps.pilah(k, tmp, buat_pdf=False)
        assert [f.name for f in h["file"]] == ["detail-resi-spesial-satuan-1-2.xlsx"]
        assert not (tmp / "SISA-SATUAN-1-2").exists() and not any((tmp / "SPESIAL-SATUAN-1-2").glob("*.pdf"))
        assert kurasi_satuan.baca() == {} and "dikurasi" not in h
        assert ps.ringkasan_json(h)["tanpa_pdf"] is True
        # sesudah mode PDF lalu hanya-Excel pada sesi sama: PDF lama hilang, penanda dicabut (tak ada pengganti)
        ps.pilah(k, tmp)
        assert kurasi_satuan.baca()
        ps.pilah(k, tmp, buat_pdf=False)
        assert kurasi_satuan.baca() == {} and not any((tmp / "SPESIAL-SATUAN-1-2").glob("*.pdf"))
        assert ps.ringkasan_json(ps.pilah(k, tmp))["tanpa_pdf"] is False
    print("  mode hanya-Excel: tanpa PDF/SISA, tanpa penanda dikurasi (dan mencabut penanda dari mode PDF sebelumnya)")


def uji_resi_dobel_dihitung_sekali():
    baris = [{"sku": "X", "resi": "R1", "pesanan": "O1", "kurir": "SPX Hemat"}] * 3
    layak, sisa = ps.kelompokkan(baris, 3)
    assert not layak and sisa == {"X": 1}
    print("  halaman dobel (resi sama) tidak menggelembungkan hitungan")


class KlienItem:
    def __init__(self, items):
        self.items = items

    def get(self, path):
        return {"items": self.items}


def _it(sku, so, resi, **extra):
    return {"item_code": sku, "salesorder_no": so, "tracking_no": resi, "shipper": "SPX Hemat",
            "wms_status": "FINISH_PICK", **extra}


def uji_awalan_ub_tidak_pernah_spesial():
    baris = [{"sku": "C225-UB11-2", "resi": f"R{i}", "pesanan": f"O{i}", "kurir": "SPX Hemat"} for i in range(5)]
    baris += [{"sku": "c225-ub10-1", "resi": f"S{i}", "pesanan": f"P{i}", "kurir": "SPX Hemat"} for i in range(3)]
    baris += [{"sku": "OK", "resi": f"T{i}", "pesanan": f"Q{i}", "kurir": "SPX Hemat"} for i in range(3)]
    layak, sisa = ps.kelompokkan(baris, 3)
    assert list(layak) == ["OK"] and not sisa, (list(layak), sisa)
    print("  SKU berawalan C225-UB (semua varian, tak peka huruf) tidak jadi spesial, tetap sisa")


def uji_bundle_multi_item_dan_qty_bukan_1_tidak_spesial():
    k = KlienItem([
        _it("A", "O1", "R1"),                                             # biasa -> boleh
        _it("B", "O2", "R2", bundle_item_id=77),                          # bundle (BD dst)
        _it("C", "O3", "R3", qty_ordered="2.0000"),                       # qty 2
        _it("D", "O4", "R4"), _it("E", "O4", "R4"),                       # 2 item satu pesanan
        _it("F", "O5", "R5", bundle_item_id=0, qty_ordered="1.0000"),     # biasa -> boleh
    ])
    boleh = {i["sku"]: i["boleh_spesial"] for i in ps.item_picklist(k, 1)}
    assert boleh == {"A": True, "B": False, "C": False, "D": False, "E": False, "F": True}, boleh
    baris = [{**i, "kurir": "SPX Hemat"} for i in ps.item_picklist(k, 1) for _ in range(3)
             if i["sku"] in ("B", "C", "D")]
    baris = [{**b, "resi": f"{b['resi']}-{n}"} for n, b in enumerate(baris)]
    layak, _ = ps.kelompokkan(baris, 1)
    assert not layak, layak
    print("  bundle, multi-item per pesanan, qty bukan 1 -> tidak spesial (tetap sisa)")


def _siapkan_sesi3(tmp: Path, k: KlienPalsu) -> KlienPalsu:
    """Sesi 3: picklist 30 berisi Y (O7/R7) -> Y gabungan sesi 1-2-3 = 3 resi (R3, R6, R7)."""
    (tmp / "3" / "SATUAN").mkdir(parents=True)
    _pdf_label(tmp / "3" / "SATUAN" / "PICK-000000030_1QTY-REGULER-2A_2026-10-13_083000.pdf", [("O7", "R7")])
    return KlienPalsu({**k.data, 30: [("Y", "O7", "R7", "FINISH_PICK")]})


def uji_pilah_lanjutan_sisa_lama_digabung_sesi_baru():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        ps.pilah(k, tmp)                                    # putaran 1: X spesial (3), Y sisa (2)
        spesial_lama = sorted((tmp / "SPESIAL-SATUAN-1-2").glob("*.pdf"))
        isi_lama = [f.read_bytes() for f in spesial_lama]
        k2 = _siapkan_sesi3(tmp, k)
        h = ps.pilah(k2, tmp, sesi=["3"], lanjutan_dari="1-2")
        assert h["sesi"] == "1-2+3" and h["sudah_spesial"] == 3, h["sesi"]   # R1, R2, R5 sudah jadi spesial
        assert set(h["layak"]) == {"Y"} and sorted(b["resi"] for b in h["layak"]["Y"]) == ["R3", "R6", "R7"], h["layak"]
        assert not h["sisa"] and not (tmp / "SISA-SATUAN-1-2+3").exists()
        pdf_y = tmp / "SPESIAL-SATUAN-1-2+3" / "SPESIAL-SATUAN_2A_Y_3resi.pdf"
        assert pdf_y.exists() and len(PdfReader(str(pdf_y)).pages) == 3
        assert (tmp / "SPESIAL-SATUAN-1-2+3" / "detail-resi-spesial-satuan-1-2+3.xlsx").exists()
        # hasil putaran 1 (yang sudah dicetak) tidak disentuh; X tidak muncul lagi di putaran 2
        assert [f.read_bytes() for f in sorted((tmp / "SPESIAL-SATUAN-1-2").glob("*.pdf"))] == isi_lama
        assert not list((tmp / "SPESIAL-SATUAN-1-2+3").glob("*_X_*"))
        # SISA-SATUAN-1-2 digantikan: tidak ada PDF di tingkat atas (tidak ikut tercetak bulk), jejak di subfolder
        sisa_lama = tmp / "SISA-SATUAN-1-2"
        assert sisa_lama.is_dir() and not list(sisa_lama.glob("*.pdf")), "folder tetap ada, PDF dipindah"
        assert len(list((sisa_lama / ps.FOLDER_SISA_TERPAKAI).glob("*.pdf"))) == 1 and h["sisa_lama_dipindah"] == 1
        assert kurasi_satuan.daftar_hasil(tmp / "3", kurasi_satuan.AWALAN_SISA) == []
        # penanda dikurasi: satuan sesi 3 -> hasil baru; satuan 1-2 tetap -> hasil putaran 1 (masih berlaku)
        aktif = kurasi_satuan.baca()
        asal = {f.parent.parent.name: aktif.get(str(f.resolve())) for f in ps.cari_pdf_satuan(tmp)}
        assert asal == {"1": "1-2", "2": "1-2", "3": "1-2+3"}, asal
        # dijalankan ulang (mis. ambang diubah) menghasilkan hal yang sama, tanpa bentrok
        h2 = ps.pilah(k2, tmp, sesi=["3"], lanjutan_dari="1-2")
        assert set(h2["layak"]) == {"Y"} and h2["sudah_spesial"] == 3
    print("  lanjutan: sisa 1-2 + satuan sesi 3 -> SPESIAL-SATUAN-1-2+3; spesial 1-2 (tercetak) tidak ikut & tidak disentuh")


def uji_pilah_lanjutan_bertingkat_dan_sisa_baru():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        ps.pilah(k, tmp)
        k2 = _siapkan_sesi3(tmp, k)
        h = ps.pilah(k2, tmp, sesi=["3"], lanjutan_dari="1-2", min_resi=4)   # Y = 3 < 4 -> tetap sisa
        assert not h["layak"] and {kel: len(bs) for kel, bs in h["sisa"].items()} == {"1QTY-REGULER-2A": 3}, h["sisa"]
        assert (tmp / "SISA-SATUAN-1-2+3" / "SISA-SATUAN_1QTY-REGULER-2A_3resi.pdf").exists()
        assert not list((tmp / "SISA-SATUAN-1-2").glob("*.pdf"))
        assert not (tmp / "SPESIAL-SATUAN-1-2+3").exists()
        # sesi 4 berisi Y lagi -> putaran 3 menggabung sisa "1-2+3" (rantai) tanpa X; Y = 4 resi
        (tmp / "4" / "SATUAN").mkdir(parents=True)
        _pdf_label(tmp / "4" / "SATUAN" / "PICK-000000040_1QTY-REGULER-2A_2026-10-13_090000.pdf", [("O8", "R8")])
        k3 = KlienPalsu({**k2.data, 40: [("Y", "O8", "R8", "FINISH_PICK")]})
        h3 = ps.pilah(k3, tmp, sesi=["4"], lanjutan_dari="1-2+3", min_resi=4)
        assert h3["sesi"] == "1-2+3+4" and h3["sudah_spesial"] == 3, (h3["sesi"], h3["sudah_spesial"])
        assert set(h3["layak"]) == {"Y"} and len(h3["layak"]["Y"]) == 4
        assert (tmp / "SPESIAL-SATUAN-1-2+3+4" / "SPESIAL-SATUAN_2A_Y_4resi.pdf").exists()
        assert not list((tmp / "SISA-SATUAN-1-2+3").glob("*.pdf"))
    print("  lanjutan bertingkat (1-2 -> 1-2+3 -> 1-2+3+4) & sisa baru jika masih di bawah ambang")


def uji_pilah_lanjutan_penolakan():
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        k = _siapkan(tmp)
        k2 = _siapkan_sesi3(tmp, k)
        for kwargs, kata in (({"sesi": ["3"], "lanjutan_dari": "1-2"}, "tidak ditemukan"),   # hasil lama belum ada
                             ({"sesi": None, "lanjutan_dari": "1-2"}, "butuh sesi baru")):
            try:
                ps.pilah(k2, tmp, **kwargs)
                assert False, kwargs
            except ps.SesiTidakValid as e:
                assert kata in str(e), e
        ps.pilah(k2, tmp, sesi=["1", "2"])
        try:   # sesi 2 sudah termasuk hasil pilah 1-2
            ps.pilah(k2, tmp, sesi=["2", "3"], lanjutan_dari="1-2")
            assert False
        except ps.SesiTidakValid as e:
            assert "sudah termasuk" in str(e), e
        # SISA lama sudah (sebagian) dicetak -> ditolak, tidak ada yang ditulis/dipindah
        sisa = next((tmp / "SISA-SATUAN-1-2").glob("*.pdf"))
        print_spesial.catat_sudah_dicetak(sisa)
        try:
            ps.pilah(k2, tmp, sesi=["3"], lanjutan_dari="1-2")
            assert False
        except ps.SesiTidakValid as e:
            assert "sudah dicetak" in str(e), e
        assert sisa.exists() and not (tmp / "SPESIAL-SATUAN-1-2+3").exists()
        # mode analisis tidak menulis apa pun
        print_spesial.lupakan_sudah_dicetak([sisa])
        h = ps.pilah(k2, tmp, sesi=["3"], lanjutan_dari="1-2", tulis=False)
        assert set(h["layak"]) == {"Y"} and sisa.exists() and not (tmp / "SPESIAL-SATUAN-1-2+3").exists()
    print("  lanjutan ditolak: hasil lama tak ada, tanpa sesi baru, sesi tumpang tindih, SISA lama sudah dicetak")


if __name__ == "__main__":
    for nama, f in list(globals().items()):
        if nama.startswith("uji_"):
            print(nama)
            f()
    print("SEMUA UJI LULUS")
