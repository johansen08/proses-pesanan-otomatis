"""Buat ulang data/template/picklist-form-kosong.xlsx dari sheet "HARI IN - FORM KOSONG" di
PICKLIST.xlsx (salinan master tim). Dijalankan MANUAL, jarang (cuma kalau tim mengubah
format/rumus form di master) - membuka file ~51 ribu baris itu makan beberapa menit.

Template hasilnya hanya berisi sheet itu dengan baris 1-6 (judul, subtotal, header, 1 baris
rumus contoh): cukup bagi rekap_master_excel.py untuk membuat file PICKLIST per sesi dalam
milidetik. Nama sheet SENGAJA dibiarkan persis seperti di master ("HARI IN - FORM KOSONG",
bukan "HARI INI") supaya kalau seluruh sheet-nya dipindah ke master tidak bentrok dengan
sheet "HARI INI" yang sudah ada.

    python src/buat_template_picklist.py [PICKLIST.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SHEET = "HARI IN - FORM KOSONG"
KOLOM_TERAKHIR = 27         # AA
BARIS_TERAKHIR = 6          # 1-5 judul/header, 6 = baris rumus contoh (rekap_master_excel.BARIS_DATA_AWAL)
TUJUAN = ROOT / "data" / "template" / "picklist-form-kosong.xlsx"


def buat(sumber: Path, tujuan: Path = TUJUAN) -> None:
    """Salin baris 1-6 ke workbook BARU (nilai + gaya yang benar-benar dipakai saja). Menghapus
    sheet/baris dari workbook master tidak cukup: stylesheet-nya tetap membawa puluhan ribu
    fill/gaya sisa 51 ribu baris, sehingga tiap wb.save() makan ~1 detik."""
    from copy import copy

    from openpyxl import Workbook

    asal = load_workbook(sumber)[SHEET]
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for baris in asal.iter_rows(min_row=1, max_row=BARIS_TERAKHIR, max_col=KOLOM_TERAKHIR):
        for c in baris:
            if not (c.has_style or c.value is not None):
                continue
            n = ws.cell(c.row, c.column, c.value)
            if c.has_style:
                n.font, n.fill, n.border = copy(c.font), copy(c.fill), copy(c.border)
                n.alignment, n.protection = copy(c.alignment), copy(c.protection)
                n.number_format = c.number_format
    for huruf, dim in asal.column_dimensions.items():
        n = ws.column_dimensions[huruf]
        n.width, n.hidden = dim.width, dim.hidden
    for r in range(1, BARIS_TERAKHIR + 1):
        ws.row_dimensions[r].height = asal.row_dimensions[r].height
    for rentang in asal.merged_cells.ranges:
        if rentang.max_row <= BARIS_TERAKHIR:
            ws.merge_cells(str(rentang))
    for cf in asal.conditional_formatting:     # mis. duplikat nomor picklist di kolom L
        for aturan in cf.rules:
            ws.conditional_formatting.add(str(cf.sqref), copy(aturan))
    ws.sheet_view.zoomScale = asal.sheet_view.zoomScale
    ws.sheet_view.showGridLines = asal.sheet_view.showGridLines
    ws.freeze_panes = "A6"              # di master "N1972" (sisa scroll), bukan header
    tujuan.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tujuan)


if __name__ == "__main__":
    buat(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "PICKLIST.xlsx")
    print(f"Template ditulis ke {TUJUAN}")
