"""Catat tiap picklist ke PICKLIST.xlsx PER SESI - file kecil di folder sesi label
(`label-pengiriman/YYYY-MM-DD/N/PICKLIST.xlsx`) yang dibuat dari template
`data/template/picklist-form-kosong.xlsx` (sheet "HARI IN - FORM KOSONG" milik tim, lengkap dengan
rumus & format). Tim memverifikasinya lalu meng-copy-paste baris-barisnya MANUAL ke file
master "PICK LIST - EXCEL ... MASTER" - program TIDAK PERNAH membuka/menulis file master
(dulu salinan ~51 ribu baris: buka ~48 detik + simpan ~67 detik per proses, lihat riwayat
insiden 2026-10-06). Template dibuat ulang lewat src/buat_template_picklist.py.

Kolom P (SCAN) diisi terpisah oleh isi_scan() setelah upload IRESIS (total resi per picklist).

Kolom yang diisi (nilainya saja, rumus bawaan template dibiarkan): F=operator (operator aktif,
lihat operator_aktif.py / data/operator.json), G=tanggal, H=jam (desimal gaya "HH.MM", mis. 18.53 = jam 18:53, BUKAN pecahan jam
sungguhan), L=nomor picklist (angka saja), M=Total Pesanan, N=Resi Keluar, T=label alur
(sama seperti kolom "SKU" di riwayat_picklist.xlsx), U=daftar no pesanan yang belum dapat
resi. Kolom rumus B/E/O/Q/W di tiap baris baru disalin dari baris rumus template (rumus
relatif, jadi tetap benar kalau ditempel ke baris lain di master).

Pewarnaan kuning: sel O (MINUS = Total Pesanan - Resi Keluar) kalau hasilnya > 0, dan seluruh
baris B:U untuk picklist GAGAL/TERHENTI serta "PICKLIST CANCEL".

"PICKLIST CANCEL" = nomor yang terlompat SELAMA proses python ini berjalan (lihat
peringatan_picklist.ambil_nomor_hilang() - nomor yang sudah terlompat SEBELUM proses mulai
tidak ditulis, hanya diperingatkan di layar: bisa jadi dibuat orang lain di PC lain).

Penulisan: file sesi kecil (puluhan baris), jadi tiap catat() langsung menambah baris & menyimpan
ATOMIK (tulis ke file sementara lalu os.replace, supaya file lama utuh kalau terputus) - tanpa
antrean. Workbook disimpan di memori per proses; kalau penyimpanan gagal (PICKLIST.xlsx sedang
dibuka di Excel) isinya tetap ada di memori dan ikut tersimpan pada catat() berikutnya atau
selesai() di akhir proses; kalau tetap gagal, dicatat lewat peringatan_gagal.
"""
from __future__ import annotations

import copy
import logging
import os
import re
import threading
import time
from datetime import date, datetime
from pathlib import Path

import operator_aktif
import peringatan_gagal

log = logging.getLogger("sku-spesial")

TEMPLATE = Path(__file__).resolve().parent.parent / "data" / "template" / "picklist-form-kosong.xlsx"
NAMA_FILE = "PICKLIST.xlsx"
NAMA_SEMENTARA = "PICKLIST.xlsx.menulis"         # lihat _simpan()
BARIS_DATA_AWAL = 6            # baris rumus contoh di template; data mulai di sini
KOLOM_F_OPR, KOLOM_G_TGL, KOLOM_H_JAM = 6, 7, 8
KOLOM_L_PICKLIST, KOLOM_M_LOLOS, KOLOM_N_PRINT, KOLOM_O_MINUS = 12, 13, 14, 15
KOLOM_P_SCAN = 16
KOLOM_T_JENIS, KOLOM_U_CATATAN = 20, 21
KOLOM_TERAKHIR = 27            # AA - kolom terjauh yang berformat di template
KOLOM_FILL_AWAL, KOLOM_FILL_AKHIR = 2, 21      # B..U - lihat baris kuning "PICKLIST CANCEL"
WARNA_KUNING = "FFFFFF00"
TEKS_TERLOMPAT = "PICKLIST CANCEL"
MAKS_COBA_GANTI_FILE = 3       # os.replace gagal (PermissionError) kalau file sedang dibuka
JEDA_COBA_GANTI_FILE_S = 0.5


class _Buku:
    """Workbook satu sesi yang sedang dikerjakan proses ini."""

    def __init__(self, file: Path, wb, ws):
        self.file, self.wb, self.ws = file, wb, ws
        self.baris_berikut = BARIS_DATA_AWAL
        while ws.cell(self.baris_berikut, KOLOM_L_PICKLIST).value is not None:
            self.baris_berikut += 1
        # rumus & gaya baris contoh - disalin ke tiap baris baru
        self.rumus = {c: ws.cell(BARIS_DATA_AWAL, c).value for c in range(1, KOLOM_TERAKHIR + 1)
                      if str(ws.cell(BARIS_DATA_AWAL, c).value or "").startswith("=")}
        self.gaya = [copy.copy(ws.cell(BARIS_DATA_AWAL, c)._style)
                     for c in range(1, KOLOM_TERAKHIR + 1)]
        self.tinggi = ws.row_dimensions[BARIS_DATA_AWAL].height
        self.kotor = False            # ada perubahan yang belum tersimpan
        self.dicatat = 0              # picklist yang ditambahkan proses ini


_buku: dict[Path, _Buku] = {}
_lock = threading.Lock()
_sudah_peringatan_template = False
_simpan_pernah_gagal = False


def _buka(folder_sesi: Path) -> _Buku | None:
    global _sudah_peringatan_template
    file = folder_sesi / NAMA_FILE
    if file in _buku:
        return _buku[file]
    from openpyxl import load_workbook

    wb = None
    if file.exists():
        try:
            wb = load_workbook(file)
        except Exception as e:     # noqa: BLE001 - file rusak jangan menggagalkan proses utama
            rusak = file.with_name(f"PICKLIST_rusak_{datetime.now():%H%M%S}.xlsx")
            log.warning("  Gagal buka %s (%s) - dipindah ke %s, dibuat baru dari template",
                        file.name, e, rusak.name)
            try:
                os.replace(file, rusak)
            except OSError:
                pass
    if wb is None:
        if not TEMPLATE.exists():
            if not _sudah_peringatan_template:
                log.warning("  Template %s tidak ditemukan - %s tidak dibuat (jalankan "
                            "src/buat_template_picklist.py)", TEMPLATE, NAMA_FILE)
                _sudah_peringatan_template = True
            return None
        wb = load_workbook(TEMPLATE)
    ws = wb.worksheets[0]
    _buku[file] = _Buku(file, wb, ws)
    return _buku[file]


# ============================================================== catat
def catat(folder_sesi: Path, baris_riwayat: dict, nomor_terlompat: list[int] | None = None) -> None:
    """Panggil di titik yang sama dengan proses_label.catat_riwayat() - 1 baris
    riwayat_picklist.xlsx = 1 baris PICKLIST.xlsx di folder sesi `folder_sesi`. `nomor_terlompat`:
    hasil peringatan_picklist.ambil_nomor_hilang() diambil SEGERA setelah picklist ini dibuat -
    tiap nomor dapat baris kuning tersendiri sebelum baris picklist ini."""
    if _angka_picklist(baris_riwayat.get("No Picklist", "")) is None:
        return
    try:
        with _lock:
            buku = _buka(Path(folder_sesi))
            if buku is None:
                return
            _tulis_baris(buku, baris_riwayat, nomor_terlompat)
            buku.kotor = True
            buku.dicatat += 1
            _simpan(buku)
    except Exception as e:     # noqa: BLE001 - jangan sampai proses picklist utama gagal
        log.warning("  Gagal catat %s ke %s: %s", baris_riwayat.get("No Picklist"), NAMA_FILE, e)


def selesai() -> None:
    """Dipanggil main.py di akhir tiap proses: coba simpan sekali lagi yang masih tertunda
    (PICKLIST.xlsx dibuka di Excel saat catat()), lalu info/peringatan jumlah yang ditulis."""
    with _lock:
        for buku in _buku.values():
            if buku.kotor:
                _simpan(buku, peringatan_terakhir=True)
            if buku.dicatat:
                log.info("%d picklist ditulis ke %s", buku.dicatat, buku.file)
                buku.dicatat = 0


# ============================================================== scan IRESIS
def isi_scan(folder_sesi: Path, total_resi: dict[int, int]) -> int:
    """Isi kolom P (SCAN) tiap baris picklist di PICKLIST.xlsx `folder_sesi` dari laporan Total
    Picklist IRESIS (`total_resi`: nomor picklist -> total resi). Dicocokkan lewat kolom L.
    File yang belum ada tidak dibuat. Mengembalikan jumlah sel yang berubah (kolom Q/CTRL
    otomatis lewat rumus template)."""
    file = Path(folder_sesi) / NAMA_FILE
    if not file.exists():
        return 0
    with _lock:
        buku = _buka(Path(folder_sesi))
        if buku is None:
            return 0
        ws, berubah = buku.ws, 0
        for r in range(BARIS_DATA_AWAL, buku.baris_berikut):
            nomor = ws.cell(r, KOLOM_L_PICKLIST).value
            baru = total_resi.get(nomor) if isinstance(nomor, int) else None
            if baru is not None and ws.cell(r, KOLOM_P_SCAN).value != baru:
                ws.cell(r, KOLOM_P_SCAN, baru)
                berubah += 1
        if berubah:
            buku.kotor = True
            _simpan(buku, peringatan_terakhir=True)
        return berubah


# ============================================================== simpan
def _ganti_file(asal: Path, tujuan: Path, coba_maks: int) -> None:
    """os.replace dengan coba ulang singkat - di Windows gagal (PermissionError) selama
    tujuan sedang dibuka Excel."""
    for coba in range(1, coba_maks + 1):
        try:
            os.replace(asal, tujuan)
            return
        except PermissionError:
            if coba == coba_maks:
                raise
            time.sleep(JEDA_COBA_GANTI_FILE_S)


def _simpan(buku: _Buku, peringatan_terakhir: bool = False) -> bool:
    """Simpan ke NAMA_SEMENTARA dulu, baru os.replace ke PICKLIST.xlsx. Sesudah satu kegagalan,
    percobaan berikutnya tidak lagi menunggu (tidak memperlambat tiap picklist selagi file
    masih dibuka di Excel)."""
    global _simpan_pernah_gagal
    sementara = buku.file.with_name(NAMA_SEMENTARA)
    try:
        buku.wb.save(sementara)
        _ganti_file(sementara, buku.file, 1 if _simpan_pernah_gagal else MAKS_COBA_GANTI_FILE)
    except OSError as e:
        _simpan_pernah_gagal = True
        sementara.unlink(missing_ok=True)
        if peringatan_terakhir:
            pesan = (f"{NAMA_FILE} sesi {buku.file.parent.name} belum tersimpan ({e}); tutup "
                     f"file itu di Excel lalu jalankan ulang, atau tulis manual dari "
                     f"riwayat_picklist.xlsx")
            log.warning("  %s", pesan)
            peringatan_gagal.catat(pesan)
        return False
    buku.kotor = False
    _simpan_pernah_gagal = False
    return True


# ============================================================== baris
def _angka_picklist(no: str) -> int | None:
    m = re.fullmatch(r"PICK-0*(\d+)", str(no).strip().upper())
    return int(m.group(1)) if m else None


def _jam_desimal(waktu: datetime) -> float:
    return float(f"{waktu.hour}.{waktu.minute:02d}")


def _isian_kuning():
    from openpyxl.styles import PatternFill

    return PatternFill(fill_type="solid", fgColor=WARNA_KUNING)


def _baris_kosong(buku: _Buku, tanggal: date) -> int:
    """Ambil baris berikutnya: salin rumus & gaya baris contoh (kecuali baris contoh itu
    sendiri), lalu isi operator & tanggal."""
    from openpyxl.formula.translate import Translator

    r = buku.baris_berikut
    buku.baris_berikut += 1
    ws = buku.ws
    if r != BARIS_DATA_AWAL:
        for c in range(1, KOLOM_TERAKHIR + 1):
            sel = ws.cell(r, c)
            sel._style = copy.copy(buku.gaya[c - 1])
            rumus = buku.rumus.get(c)
            if rumus:
                sel.value = Translator(rumus, origin=f"{ws.cell(BARIS_DATA_AWAL, c).coordinate}"
                                       ).translate_formula(sel.coordinate)
        if buku.tinggi:
            ws.row_dimensions[r].height = buku.tinggi
    ws.cell(r, KOLOM_F_OPR, operator_aktif.ambil_aktif())
    ws.cell(r, KOLOM_G_TGL, datetime.combine(tanggal, datetime.min.time()))
    return r


def _tandai_kuning(ws, baris: int, kolom_awal: int = KOLOM_FILL_AWAL,
                   kolom_akhir: int = KOLOM_FILL_AKHIR) -> None:
    isian = _isian_kuning()
    for c in range(kolom_awal, kolom_akhir + 1):
        ws.cell(baris, c).fill = isian


def _tulis_baris(buku: _Buku, baris_riwayat: dict, nomor_terlompat: list[int] | None) -> None:
    """Tambahkan 1 picklist (+ baris "PICKLIST CANCEL" untuk tiap nomor terlompat sebelumnya)."""
    ws = buku.ws
    n = _angka_picklist(baris_riwayat["No Picklist"])
    waktu = (datetime.strptime(baris_riwayat["Waktu"], "%d-%m-%Y %H:%M")
             if baris_riwayat.get("Waktu") else datetime.now())
    for hilang in nomor_terlompat or []:
        r = _baris_kosong(buku, waktu.date())
        ws.cell(r, KOLOM_L_PICKLIST, hilang)
        ws.cell(r, KOLOM_T_JENIS, TEKS_TERLOMPAT)
        _tandai_kuning(ws, r)
    r = _baris_kosong(buku, waktu.date())
    ws.cell(r, KOLOM_H_JAM, _jam_desimal(waktu))
    ws.cell(r, KOLOM_L_PICKLIST, n)
    total, resi = baris_riwayat.get("Total Pesanan"), baris_riwayat.get("Resi Keluar")
    if total is not None:
        ws.cell(r, KOLOM_M_LOLOS, total)
    if resi is not None:
        ws.cell(r, KOLOM_N_PRINT, resi)
    if baris_riwayat.get("SKU") is not None:
        ws.cell(r, KOLOM_T_JENIS, baris_riwayat["SKU"])
    if baris_riwayat.get("Tanpa Resi"):
        ws.cell(r, KOLOM_U_CATATAN, ", ".join(baris_riwayat["Tanpa Resi"]))
    if str(baris_riwayat.get("Catatan", "")).startswith(("GAGAL", "TERHENTI")):
        _tandai_kuning(ws, r)
    elif isinstance(total, (int, float)) and isinstance(resi, (int, float)) and total - resi > 0:
        _tandai_kuning(ws, r, KOLOM_O_MINUS, KOLOM_O_MINUS)      # MINUS > 0
